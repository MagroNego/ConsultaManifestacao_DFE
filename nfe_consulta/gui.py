"""Interface grafica Windows para o CLI de consulta de manifestacoes."""

import os
import queue
import re
import sqlite3
import subprocess
import sys
import threading
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, simpledialog, ttk

from openpyxl import load_workbook

from nfe_consulta import __version__
from nfe_consulta.assistente import CNPJ_YAB, NOME_PLANILHA, UF_YAB
from nfe_consulta.status import consultar_status
from nfe_consulta.seguranca_banco import criptografado


AZUL = "#37326F"
AZUL_CLARO = "#EEECFC"
FUNDO = "#F6F7FC"
BRANCO = "#FFFFFF"
TEXTO = "#252640"
CINZA = "#626781"
VERDE = "#276A60"
ROXO = "#6155D9"
BORDA = "#E3E5F0"


def caminhos_iniciais(pasta: Path, downloads: Path) -> tuple[Path, Path, Path]:
    chaves = pasta / "entrada" / "CHAVES.txt"
    if not chaves.is_file():
        chaves = downloads / "CHAVES.txt"
    banco = pasta / "dados" / "nfe_manifestacoes.db"
    for candidato in (pasta / "dados" / "nfe_manifestacoes_seguro.db",
                      downloads / "nfe_manifestacoes_seguro.db", banco,
                      downloads / "nfe_manifestacoes.db"):
        if candidato.exists():
            banco = candidato
            break
    return chaves, banco, pasta / "saidas" / NOME_PLANILHA


def montar_comando(chaves: Path, banco: Path, saida: Path, sincronizar: bool,
                   max_lotes: int = 50) -> list[str]:
    comando = [
        sys.executable, "-u", "-m", "nfe_consulta.cli", "--cnpj", CNPJ_YAB,
        "--lote", str(chaves), "--xlsx", str(saida), "--banco", str(banco),
    ]
    if sincronizar:
        if not 1 <= max_lotes <= 500:
            raise ValueError("Lotes por execucao devem estar entre 1 e 500")
        comando += ["--manifestacoes", "--uf", UF_YAB, "--max-lotes", str(max_lotes)]
    return comando


def ler_previa(caminho: Path) -> tuple[list[tuple[str, ...]], str]:
    """Le somente as colunas necessarias da planilha gerada pelo aplicativo."""
    arquivo = load_workbook(caminho, read_only=True, data_only=True)
    try:
        aba = arquivo.active
        resumo = str(aba["A2"].value or "")
        linhas = []
        for chave, numero, serie, evento, data, protocolo, _, _, historico, erro in aba.iter_rows(
            min_row=5, min_col=1, max_col=10, values_only=True
        ):
            if chave is not None:
                linhas.append(tuple(str(valor) if valor is not None else "" for valor in (
                    chave, numero, serie, evento, data, protocolo, historico, erro,
                )))
        return linhas, resumo
    finally:
        arquivo.close()


def resumir_status(texto: str) -> tuple[str, str, str, str]:
    """Extrai apenas os dados que já constam no status local, sem consultar a SEFAZ."""
    nsus = re.search(r"ultNSU salvo: (\d+) \| maxNSU conhecido: (\d+)", texto)
    resposta = re.search(r"Ultima resposta da SEFAZ salva: (\d{2}/\d{2}/\d{4} \d{2}:\d{2})", texto)
    if "Não foi possível" in texto or "nao encontrado" in texto.lower():
        condicao = "Banco indisponível"
    elif "Pausa por" in texto:
        condicao = "Em pausa"
    elif "Sincronizacao parcial" in texto:
        condicao = "Parcial"
    elif nsus:
        condicao = "Fila percorrida"
    else:
        condicao = "Sem histórico"
    return (resposta.group(1) if resposta else "Sem registro",
            nsus.group(1) if nsus else "—",
            nsus.group(2) if nsus else "—", condicao)


class AplicativoManifestacoes(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.pasta = Path(__file__).resolve().parent.parent
        chaves, banco, saida = caminhos_iniciais(self.pasta, Path.home() / "Downloads")
        self.chaves = tk.StringVar(value=str(chaves))
        self.banco = tk.StringVar(value=str(banco))
        self.saida = tk.StringVar(value=str(saida))
        self.max_lotes = tk.StringVar(value="50")
        self.pesquisa = tk.StringVar()
        self.resumo = tk.StringVar(value="Nenhuma planilha gerada nesta sessão")
        self.contagem = tk.StringVar(value="0 resultados")
        self.situacao = tk.StringVar(value="Pronto para consultar o banco local")
        self.fila: queue.Queue = queue.Queue()
        self.linhas: list[tuple[str, ...]] = []
        self.executando = False
        self._senha_banco: str | None = None
        self._caminho_senha: str | None = None
        self.filtro_agendado: str | None = None
        self.title("Consulta de Manifestação | YAB")
        self.geometry("1200x790")
        self.minsize(980, 650)
        self.status_resposta = tk.StringVar(value="Sem registro")
        self.status_ult_nsu = tk.StringVar(value="—")
        self.status_max_nsu = tk.StringVar(value="—")
        self.status_condicao = tk.StringVar(value="Aguardando leitura")
        self.configure(bg=FUNDO)
        self._estilo()
        self._montar()
        self._atualizar_status()
        self.after(100, self._processar_fila)
        self.protocol("WM_DELETE_WINDOW", self._fechar)

    def _estilo(self) -> None:
        estilo = ttk.Style(self)
        if "clam" in estilo.theme_names():
            estilo.theme_use("clam")
        estilo.configure("YAB.Treeview", background=BRANCO, foreground=TEXTO,
                         fieldbackground=BRANCO, rowheight=32, font=("Segoe UI", 10), borderwidth=0)
        estilo.configure("YAB.Treeview.Heading", background="#F0F0FA", foreground=AZUL,
                         font=("Segoe UI", 10, "bold"), padding=(9, 10), relief="flat")
        estilo.map("YAB.Treeview", background=[("selected", "#E2DFFD")],
                   foreground=[("selected", AZUL)])
        estilo.configure("YAB.Horizontal.TProgressbar", troughcolor="#E4E4F2", background=ROXO)

    def _rotulo(self, pai, texto: str = "", tamanho: int = 10, cor: str = TEXTO,
                negrito: bool = False, fundo: str = BRANCO, **kw) -> tk.Label:
        return tk.Label(pai, text=texto, bg=fundo, fg=cor,
                        font=("Segoe UI", tamanho, "bold" if negrito else "normal"), **kw)

    def _botao(self, pai, texto: str, acao, fundo: str, cor: str = BRANCO,
               **kw) -> tk.Button:
        opcoes = dict(text=texto, command=acao, bg=fundo, fg=cor,
                      activebackground=fundo, activeforeground=cor, relief="flat",
                      borderwidth=0, cursor="hand2", font=("Segoe UI", 10, "bold"),
                      padx=18, pady=11)
        opcoes.update(kw)
        return tk.Button(pai, **opcoes)

    def _montar(self) -> None:
        cabecalho = tk.Frame(self, bg=BRANCO, height=83)
        cabecalho.pack(fill="x")
        cabecalho.pack_propagate(False)
        tk.Frame(cabecalho, bg=ROXO, width=6).pack(side="left", fill="y")
        titulos = tk.Frame(cabecalho, bg=BRANCO)
        titulos.pack(side="left", padx=24, pady=13)
        self._rotulo(titulos, "Consulta de Manifestação", 20, AZUL, True).pack(anchor="w")
        self._rotulo(titulos, f"NF-e emitidas · YAB", 10, CINZA).pack(anchor="w")
        self._botao(cabecalho, "Manual", self._abrir_manual, AZUL_CLARO, AZUL,
                    pady=6).pack(side="right", padx=(0, 25))
        self._rotulo(cabecalho, f"v{__version__}", 9, CINZA).pack(side="right", padx=18)

        area = tk.Frame(self, bg=FUNDO)
        area.pack(fill="both", expand=True, padx=24, pady=(16, 12))
        area.grid_columnconfigure(0, weight=1)
        area.grid_rowconfigure(3, weight=1)

        status = tk.Frame(area, bg=BRANCO, highlightbackground=BORDA, highlightthickness=1)
        status.grid(row=0, column=0, sticky="ew", pady=(0, 12))
        titulo_status = tk.Frame(status, bg=BRANCO)
        titulo_status.pack(fill="x", padx=17, pady=(10, 2))
        self._rotulo(titulo_status, "Situação do banco", 11, AZUL, True).pack(side="left")
        self._botao(titulo_status, "Ver status completo", self._alternar_status,
                    AZUL_CLARO, AZUL, pady=5).pack(side="right", padx=(8, 0))
        self._botao(titulo_status, "Recarregar", self._atualizar_status,
                    AZUL_CLARO, AZUL, pady=5).pack(side="right")
        dados_status = tk.Frame(status, bg=BRANCO)
        dados_status.pack(fill="x", padx=17, pady=(2, 4))
        self._rotulo(dados_status, "Última resposta:", 9, CINZA).pack(side="left")
        self._rotulo(dados_status, textvariable=self.status_resposta, tamanho=9, negrito=True).pack(
            side="left", padx=(5, 22))
        self._rotulo(dados_status, textvariable=self.status_condicao, tamanho=9,
                     cor=ROXO, negrito=True).pack(side="left")
        nsus_status = tk.Frame(status, bg=BRANCO)
        nsus_status.pack(fill="x", padx=17, pady=(0, 3))
        self._rotulo(nsus_status, "NSU salvo / máximo:", 9, CINZA).pack(side="left")
        self._rotulo(nsus_status, textvariable=self.status_ult_nsu, tamanho=9, negrito=True).pack(
            side="left", padx=(5, 3))
        self._rotulo(nsus_status, "/", 9, CINZA).pack(side="left")
        self._rotulo(nsus_status, textvariable=self.status_max_nsu, tamanho=9, negrito=True).pack(
            side="left", padx=(3, 18))
        self._rotulo(status, "Dados locais: novos eventos podem ter surgido após a última resposta.",
                     9, CINZA).pack(anchor="w", padx=17, pady=(0, 10))
        self.status_texto = tk.Text(status, height=5, wrap="word", bg=BRANCO, fg=TEXTO,
                                    relief="flat", borderwidth=0, font=("Segoe UI", 10),
                                    padx=12, pady=8, cursor="arrow")
        self.status_texto.pack(fill="x", padx=12, pady=(0, 8))
        self.status_texto.pack_forget()
        self.status_visivel = False
        self.status_texto.configure(state="disabled")

        entradas = tk.Frame(area, bg=BRANCO, highlightbackground=BORDA, highlightthickness=1)
        entradas.grid(row=1, column=0, sticky="ew", pady=(0, 11))
        entradas.grid_columnconfigure(1, weight=1)
        self._rotulo(entradas, "Arquivos da consulta", 11, AZUL, True).grid(
            row=0, column=0, columnspan=3, sticky="w", padx=19, pady=(13, 8))
        self._campo(entradas, 1, "Chaves.txt", self.chaves, self._escolher_chaves)
        self._campo(entradas, 2, "Banco", self.banco, self._escolher_banco)
        self._campo(entradas, 3, "Saída Excel", self.saida, self._escolher_saida)

        acoes = tk.Frame(area, bg=FUNDO)
        acoes.grid(row=2, column=0, sticky="ew", pady=(0, 11))
        self.botao_local = self._botao(acoes, "Gerar Excel com banco local",
                                      lambda: self._executar(False), VERDE)
        self.botao_local.pack(side="left")
        self.botao_sefaz = self._botao(acoes, "Consultar SEFAZ + Gerar Excel",
                                      lambda: self._executar(True), AZUL)
        self.botao_sefaz.pack(side="left", padx=(11, 0))
        self._rotulo(acoes, "Máx. lotes:", 9, CINZA, fundo=FUNDO).pack(side="left", padx=(20, 5))
        tk.Spinbox(acoes, from_=1, to=500, increment=10, width=5,
                   textvariable=self.max_lotes, font=("Segoe UI", 10), relief="solid",
                   borderwidth=1).pack(side="left")

        resultados = tk.Frame(area, bg=BRANCO, highlightbackground=BORDA, highlightthickness=1)
        resultados.grid(row=3, column=0, sticky="nsew", pady=(0, 12))
        resultados.grid_columnconfigure(0, weight=1)
        resultados.grid_rowconfigure(1, weight=1)
        linha_titulo = tk.Frame(resultados, bg=BRANCO)
        linha_titulo.grid(row=0, column=0, sticky="ew", padx=18, pady=(11, 9))
        self._rotulo(linha_titulo, "Resultados", 11, AZUL, True).pack(side="left")
        self._rotulo(linha_titulo, textvariable=self.contagem, cor=CINZA).pack(side="left", padx=13)
        self._botao(linha_titulo, "Abrir Excel", self._abrir_excel, AZUL_CLARO,
                    AZUL, pady=5).pack(side="right", padx=(8, 0))
        busca = tk.Entry(linha_titulo, textvariable=self.pesquisa, width=29,
                         font=("Segoe UI", 10), relief="solid", bd=1)
        busca.pack(side="right", padx=(8, 0), ipady=6)
        self._rotulo(linha_titulo, "Buscar nota ou chave:", 9, CINZA).pack(side="right")
        self.pesquisa.trace_add("write", self._agendar_filtro)
        self.campo_busca = busca

        tabela = tk.Frame(resultados, bg=BRANCO)
        tabela.grid(row=1, column=0, sticky="nsew", padx=17)
        tabela.grid_columnconfigure(0, weight=1)
        tabela.grid_rowconfigure(0, weight=1)
        colunas = ("numero", "serie", "evento", "data", "chave", "protocolo")
        self.arvore = ttk.Treeview(tabela, columns=colunas, show="headings",
                                    style="YAB.Treeview", selectmode="browse")
        for chave, titulo, largura in (
            ("numero", "NOTA", 92), ("serie", "SÉRIE", 75),
            ("evento", "MANIFESTAÇÃO", 220), ("data", "DATA DO EVENTO", 205),
            ("chave", "CHAVE DE ACESSO", 400), ("protocolo", "PROTOCOLO", 175),
        ):
            self.arvore.heading(chave, text=titulo)
            self.arvore.column(chave, width=largura, minwidth=largura // 2, stretch=chave in ("evento", "chave"))
        self.arvore.grid(row=0, column=0, sticky="nsew")
        barra_vertical = ttk.Scrollbar(tabela, orient="vertical", command=self.arvore.yview)
        barra_vertical.grid(row=0, column=1, sticky="ns")
        barra_horizontal = ttk.Scrollbar(tabela, orient="horizontal", command=self.arvore.xview)
        barra_horizontal.grid(row=1, column=0, sticky="ew")
        self.arvore.bind("<<TreeviewSelect>>", self._detalhes_selecao)
        self.arvore.configure(yscrollcommand=barra_vertical.set, xscrollcommand=barra_horizontal.set)
        self.arvore.tag_configure("alternada", background="#F8F8FD")

        detalhes = tk.Frame(resultados, bg=BRANCO)
        detalhes.grid(row=2, column=0, sticky="ew", padx=18, pady=(8, 4))
        self._rotulo(detalhes, "Histórico da nota", 9, AZUL, True).pack(anchor="w")
        self.detalhes = self._rotulo(detalhes, "Selecione uma linha para ver os eventos.", 9, CINZA,
                                     anchor="w", justify="left", wraplength=1050)
        self.detalhes.pack(fill="x", pady=(4, 7))

        rodape = tk.Frame(area, bg=FUNDO)
        rodape.grid(row=4, column=0, sticky="ew")
        self.progresso = ttk.Progressbar(rodape, mode="indeterminate", length=130,
                                         style="YAB.Horizontal.TProgressbar")
        self.progresso.pack(side="left", padx=(0, 12))
        self._rotulo(rodape, textvariable=self.situacao, cor=CINZA, fundo=FUNDO).pack(side="left")
        self._botao(rodape, "Mostrar detalhes", self._alternar_log, AZUL_CLARO,
                    AZUL, pady=5).pack(side="right")
        self._rotulo(rodape, textvariable=self.resumo, tamanho=9, cor=CINZA,
                     fundo=FUNDO).pack(side="right", padx=12)

        self.log = tk.Text(area, height=6, bg="#EEECFC", fg=AZUL,
                            relief="flat", font=("Consolas", 9), wrap="word",
                            state="disabled")
        self.log_visivel = False

    def _campo(self, pai, linha: int, titulo: str, variavel: tk.StringVar, selecionar) -> None:
        self._rotulo(pai, titulo, 9, CINZA).grid(row=linha, column=0, sticky="w", padx=(19, 17), pady=5)
        tk.Entry(pai, textvariable=variavel, font=("Segoe UI", 10), fg=TEXTO,
                 bg="#F9FBFD", relief="solid", borderwidth=1).grid(
            row=linha, column=1, sticky="ew", ipady=6, pady=5)
        self._botao(pai, "Escolher...", selecionar, AZUL_CLARO, AZUL,
                    pady=5).grid(row=linha, column=2, padx=(12, 19), pady=5)
        if linha == 3:
            tk.Frame(pai, height=7, bg=BRANCO).grid(row=4, column=0)

    def _escolher_chaves(self) -> None:
        caminho = filedialog.askopenfilename(parent=self, title="Escolher TXT de chaves",
                    initialdir=str(Path(self.chaves.get()).parent), filetypes=[("Texto", "*.txt"), ("Todos", "*.*")])
        if caminho:
            self.chaves.set(caminho)

    def _escolher_banco(self) -> None:
        caminho = filedialog.askopenfilename(parent=self, title="Escolher banco SQLite",
                    initialdir=str(Path(self.banco.get()).parent), filetypes=[("SQLite", "*.db"), ("Todos", "*.*")])
        if caminho:
            self.banco.set(caminho)
            self._senha_banco = None
            self._caminho_senha = None
            self._atualizar_status()

    def _escolher_saida(self) -> None:
        caminho = filedialog.asksaveasfilename(parent=self, title="Salvar planilha em",
                    initialdir=str(Path(self.saida.get()).parent), initialfile=NOME_PLANILHA,
                    defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if caminho:
            self.saida.set(caminho)

    def _atualizar_status(self) -> None:
        try:
            senha = self._senha_para_banco()
            estado = consultar_status(self.banco.get(), CNPJ_YAB, senha=senha)
            linhas = [l for l in estado.splitlines() if not l.startswith(("Banco:", "CNPJ:"))]
            texto = "\n".join(linhas)
        except (OSError, ValueError, RuntimeError, sqlite3.DatabaseError) as exc:
            if "Senha incorreta" in str(exc):
                self._senha_banco = None
            texto = f"Não foi possível ler o banco: {exc}"
        resposta, ult_nsu, max_nsu, condicao = resumir_status(texto)
        self.status_resposta.set(resposta)
        self.status_ult_nsu.set(ult_nsu)
        self.status_max_nsu.set(max_nsu)
        self.status_condicao.set(condicao)
        self.status_texto.configure(state="normal")
        self.status_texto.delete("1.0", "end")
        self.status_texto.insert("1.0", texto)
        self.status_texto.configure(state="disabled")

    def _senha_para_banco(self) -> str | None:
        caminho = self.banco.get()
        if not criptografado(caminho):
            return None
        if self._caminho_senha != caminho or not self._senha_banco:
            self._senha_banco = simpledialog.askstring(
                "Banco protegido", "Senha do banco criptografado:", show="*", parent=self)
            self._caminho_senha = caminho
        return self._senha_banco

    def _alternar_status(self) -> None:
        self.status_visivel = not self.status_visivel
        if self.status_visivel:
            self.status_texto.pack(fill="x", padx=12, pady=(0, 8))
        else:
            self.status_texto.pack_forget()

    def _abrir_manual(self) -> None:
        caminho = self.pasta / "MANUAL_DE_USO.md"
        if not caminho.is_file():
            messagebox.showinfo("Manual", f"Manual não encontrado em:\n{caminho}", parent=self)
            return
        try:
            os.startfile(str(caminho))
        except (AttributeError, OSError):
            messagebox.showinfo("Manual", f"Abra o arquivo:\n{caminho}", parent=self)

    def _executar(self, sincronizar: bool) -> None:
        if self.executando:
            return
        chaves = Path(self.chaves.get().strip().strip('"')).expanduser()
        banco = Path(self.banco.get().strip().strip('"')).expanduser()
        saida = Path(self.saida.get().strip().strip('"')).expanduser()
        if not chaves.is_file():
            messagebox.showerror("TXT não encontrado", f"Escolha o arquivo de chaves:\n{chaves}", parent=self)
            return
        if saida.suffix.lower() != ".xlsx":
            messagebox.showerror("Extensão incorreta", "A planilha deve terminar em .xlsx.", parent=self)
            return
        if not sincronizar and not banco.is_file():
            messagebox.showerror("Banco não encontrado", f"Escolha o banco já sincronizado:\n{banco}", parent=self)
            return
        if sincronizar:
            try:
                limite = int(self.max_lotes.get())
                if not 1 <= limite <= 500:
                    raise ValueError
            except ValueError:
                messagebox.showerror("Limite inválido", "Use um número de lotes entre 1 e 500.", parent=self)
                return
            if not banco.is_file() and not messagebox.askyesno(
                "Banco novo", "Este banco ainda não existe. Um banco novo começa sem o NSU salvo.\n\n"
                "Confirma a criação de um novo banco de consulta?", parent=self,
            ):
                return
            if not messagebox.askyesno(
                "Consulta à SEFAZ", "Esta opção usa o certificado do Windows e consulta o serviço "
                "NFeDistribuicaoDFe para o CNPJ da YAB.\n\n"
                "Aguarde a pausa para evitar evento de rejeição 656 "
                "Deseja consultar a SEFAZ agora?", parent=self,
            ):
                return
        else:
            limite = 50
        comando = montar_comando(chaves, banco, saida, sincronizar, limite)
        try:
            senha = self._senha_para_banco()
            if criptografado(banco) and not senha:
                return
        except OSError as exc:
            messagebox.showerror("Banco indisponível", str(exc), parent=self)
            return
        self.executando = True
        self.botao_local.configure(state="disabled")
        self.botao_sefaz.configure(state="disabled")
        self.progresso.start(13)
        self.situacao.set("Consultando a SEFAZ..." if sincronizar else "Gerando planilha com o banco local...")
        self._registrar("Sincronizando com a SEFAZ" if sincronizar else "Consulta somente local")
        threading.Thread(target=self._trabalhar, args=(comando, saida, senha), daemon=True).start()

    def _trabalhar(self, comando: list[str], saida: Path, senha: str | None = None) -> None:
        try:
            flags = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
            processo = subprocess.Popen(comando, cwd=self.pasta,
                        stdin=subprocess.PIPE if senha else subprocess.DEVNULL,
                        stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                        text=True, encoding="utf-8", errors="replace", bufsize=1,
                        creationflags=flags)
            if senha and processo.stdin:
                processo.stdin.write(senha + "\n")
                processo.stdin.close()
            assert processo.stdout is not None
            with processo.stdout:
                for linha in processo.stdout:
                    self.fila.put(("log", linha.rstrip("\r\n")))
            retorno = processo.wait()
            previa = ler_previa(saida) if retorno == 0 else None
            self.fila.put(("fim", retorno, previa, str(saida)))
        except Exception as exc:
            self.fila.put(("falha", str(exc)))

    def _processar_fila(self) -> None:
        try:
            while True:
                evento = self.fila.get_nowait()
                if evento[0] == "log":
                    self._registrar(evento[1])
                elif evento[0] == "fim":
                    self._finalizar()
                    codigo, previa, destino = evento[1], evento[2], evento[3]
                    if codigo == 0 and previa is not None:
                        self.linhas, resumo = previa
                        self.resumo.set(resumo[:55])
                        self._mostrar_linhas()
                        self.ultimo_arquivo = Path(destino)
                        self.situacao.set("Planilha gerada. Use 'Abrir Excel' para visualizar.")
                    else:
                        self.resumo.set("Nenhum resultado novo; a planilha anterior foi preservada")
                        self.situacao.set("Consulta não concluída. Abra 'Mostrar detalhes'.")
                        self._alternar_log(True)
                    self._atualizar_status()
                elif evento[0] == "falha":
                    self._finalizar()
                    self._registrar(f"Falha: {evento[1]}")
                    self.situacao.set("Falha ao executar. Abra 'Mostrar detalhes'.")
                    self._alternar_log(True)
                    self._atualizar_status()
        except queue.Empty:
            pass
        self.after(100, self._processar_fila)

    def _finalizar(self) -> None:
        self.executando = False
        self.progresso.stop()
        self.botao_local.configure(state="normal")
        self.botao_sefaz.configure(state="normal")

    def _registrar(self, linha: str) -> None:
        self.log.configure(state="normal")
        self.log.insert("end", linha + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def _alternar_log(self, mostrar: bool | None = None) -> None:
        novo = not self.log_visivel if mostrar is None else mostrar
        if novo == self.log_visivel:
            return
        self.log_visivel = novo
        if novo:
            self.log.grid(row=5, column=0, sticky="ew", pady=(7, 0))
        else:
            self.log.grid_remove()

    def _agendar_filtro(self, *_args) -> None:
        if self.filtro_agendado:
            self.after_cancel(self.filtro_agendado)
        self.filtro_agendado = self.after(180, self._mostrar_linhas)

    def _mostrar_linhas(self) -> None:
        self.filtro_agendado = None
        self.arvore.delete(*self.arvore.get_children())
        termo = self.pesquisa.get().strip().lower()
        correspondentes = [linha for linha in self.linhas if
                          not termo or termo in linha[0].lower() or termo in linha[1].lower()
                          or termo in linha[3].lower()]
        for i, linha in enumerate(correspondentes[:1000]):
            chave, numero, serie, evento, data, protocolo, _, erro = linha
            self.arvore.insert("", "end", iid=str(i), values=(
                numero, serie, erro or evento, data, chave, protocolo),
                tags=("alternada",) if i % 2 else ())
        self.contagem.set(f"{len(correspondentes)} resultados" +
                          (" · exibindo 1.000; use a busca" if len(correspondentes) > 1000 else ""))
        self.detalhes.configure(text="Selecione uma linha para ver os eventos.")
        self._linhas_visiveis = correspondentes[:1000]

    def _detalhes_selecao(self, _evento=None) -> None:
        selecao = self.arvore.selection()
        if not selecao or not hasattr(self, "_linhas_visiveis"):
            return
        historico = self._linhas_visiveis[int(selecao[0])][6]
        erro = self._linhas_visiveis[int(selecao[0])][7]
        self.detalhes.configure(text=erro or historico or "Nenhum evento localizado no histórico consultado.")

    def _abrir_excel(self) -> None:
        caminho = getattr(self, "ultimo_arquivo", Path(self.saida.get()).expanduser())
        if not caminho.is_file():
            messagebox.showinfo("Planilha não encontrada", "Gere a planilha antes de abri-la.", parent=self)
            return
        try:
            os.startfile(str(caminho))
        except (AttributeError, OSError) as exc:
            messagebox.showinfo("Local da planilha", f"Abra o arquivo no Excel:\n{caminho}\n{exc}", parent=self)

    def _fechar(self) -> None:
        if self.executando:
            messagebox.showinfo("Consulta em andamento", "Aguarde o término da consulta antes de fechar.", parent=self)
        else:
            self.destroy()


def main() -> None:
    aplicativo = AplicativoManifestacoes()
    aplicativo.mainloop()


if __name__ == "__main__":
    main()
