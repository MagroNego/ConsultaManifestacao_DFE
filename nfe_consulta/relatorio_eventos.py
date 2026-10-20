"""Exportação de consultas de eventos para Excel."""

from __future__ import annotations

from nfe_consulta.export_control import cancellable_rows, check_export_cancelled
import os
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from nfe_consulta.cte import spreadsheet_label
from nfe_consulta.xlsx_style import configure_sheet, style_cell, text_row_height
from nfe_consulta.web.consulta_local import EventoNota, FiltrosEventos, TIPOS_MANIFESTACAO


HEADERS = [
    "Número NF",
    "Série",
    "Chave NF-e",
    "Código",
    "Manifestação",
    "Data do evento",
    "Protocolo",
    "NSU",
    "Recebido no banco (UTC)",
    "Emitente",
    "Situação",
    "CT-e vinculado",
]


def _excel_datetime(valor: str) -> datetime | str:
    valor = (valor or "").strip()
    if not valor:
        return ""
    try:
        instante = datetime.fromisoformat(valor.replace("Z", "+00:00"))
    except ValueError:
        return valor

    return instante.replace(tzinfo=None)


def _resumo_filtros(filtros: FiltrosEventos) -> str:
    itens: list[str] = []
    if filtros.data_inicial or filtros.data_final:
        inicio = filtros.data_inicial.strftime("%d/%m/%Y") if filtros.data_inicial else "início"
        fim = filtros.data_final.strftime("%d/%m/%Y") if filtros.data_final else "hoje"
        itens.append(f"Período: {inicio} a {fim}")
    if filtros.numero is not None:
        itens.append(f"NF: {filtros.numero}")
    if filtros.serie is not None:
        itens.append(f"Série: {filtros.serie}")
    if filtros.chave:
        itens.append(f"Chave: {filtros.chave}")
    if filtros.codigo:
        itens.append(TIPOS_MANIFESTACAO.get(filtros.codigo, filtros.codigo))
    if filtros.situacao:
        itens.append('Situação: ' + filtros.situacao.capitalize())
    return "  |  ".join(itens) if itens else "Todos os eventos do banco"


def gravar_eventos_xlsx(
    caminho: str | Path,
    eventos: tuple[EventoNota, ...] | list[EventoNota],
    filtros: FiltrosEventos,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Manifestacoes"
    ws.sheet_view.zoomScale = 85
    ws.freeze_panes = "A2"

    for coluna, titulo in enumerate(HEADERS, 1):
        ws.cell(1, coluna, titulo)

    for indice, evento in enumerate(cancellable_rows(eventos), 2):
        valores = [
            evento.numero,
            evento.serie,
            evento.chave,
            evento.codigo,
            evento.descricao,
            _excel_datetime(evento.data),
            evento.protocolo,
            evento.nsu,
            _excel_datetime(evento.recebido_em),
            evento.emitente,
            evento.situacao,
            spreadsheet_label(evento.ctes),
        ]

        for coluna, valor in enumerate(valores, 1):
            celula = ws.cell(indice, coluna, valor)
            if isinstance(valor, str):
                celula.data_type = "s"
            style_cell(celula, alternate=indice % 2 == 0, wrap=coluna in (5, 10), center=coluna in (2, 4))

        for coluna in (2, 3, 4, 7, 8):
            ws.cell(indice, coluna).number_format = "@"
        ws.cell(indice, 1).number_format = "0"
        for coluna in (6, 9):
            if isinstance(ws.cell(indice, coluna).value, datetime):
                ws.cell(indice, coluna).number_format = "dd/mm/yyyy hh:mm:ss"

        ws.row_dimensions[indice].height = max(text_row_height(evento.descricao, 34), text_row_height(evento.emitente, 44))

    configure_sheet(ws, header_row=1, widths=[14, 10, 51, 13, 34, 24, 23, 19, 25, 44, 19, 23])
    ws.auto_filter.ref = f"A1:L{max(1, 1 + len(eventos))}"

    destino = Path(caminho)
    destino.parent.mkdir(parents=True, exist_ok=True)
    descritor, temporario = tempfile.mkstemp(
        prefix=".nfe-eventos-",
        suffix=".xlsx",
        dir=destino.parent,
    )
    os.close(descritor)
    try:
        check_export_cancelled()
        wb.save(temporario)
        check_export_cancelled()
        os.replace(temporario, destino)
    finally:
        Path(temporario).unlink(missing_ok=True)
