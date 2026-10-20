"""Planilha de consulta pronta para abrir no Excel, mantendo chaves como texto."""

from nfe_consulta.export_control import cancellable_rows, check_export_cancelled
import os
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from nfe_consulta.cte import spreadsheet_label
from nfe_consulta.xlsx_style import configure_sheet, style_cell, text_row_height
from nfe_consulta.cancelamento import status_label
from nfe_consulta.validacao import numero_e_serie


HEADERS = [
    "Chave NF-e", "Número", "Série", "Manifestação", "Data do evento",
    "Protocolo", "Código", "Eventos", "Histórico", "Erro", "Situação", "CT-e vinculado",
]
DESCRICOES = {
    "210200": "Confirmação da Operação",
    "210210": "Ciência da Operação",
    "210220": "Desconhecimento da Operação",
    "210240": "Operação não Realizada",
}


def _data_exibicao(iso: str) -> str:
    if not iso:
        return ""
    try:
        data = datetime.fromisoformat(iso)
    except ValueError:
        return iso
    fuso = data.strftime("%z")
    sufixo = f" UTC{fuso[:3]}:{fuso[3:]}" if fuso else ""
    return data.strftime("%d/%m/%Y %H:%M") + sufixo


def gravar_xlsx(caminho: str, resultados: list, cobertura: str, *, numeros_solicitados: dict[int, int] | None = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Consulta_Manifestacao_YAB"
    ws.sheet_view.zoomScale = 85
    ws.freeze_panes = "D5"

    ws["A1"] = "Consulta de manifestações"

    encontradas = sum(bool(r.manifestacoes) for r in resultados)
    ws["A2"] = f"{len(resultados)} chaves | {encontradas} com evento | {len(resultados)-encontradas} sem evento localizado"
    ws.row_dimensions[3].height = 8

    for coluna, titulo in enumerate(HEADERS, 1):
        ws.cell(4, coluna, titulo)
    for indice, resultado in enumerate(cancellable_rows(resultados), 5):
        numero, serie = numero_e_serie(resultado.chave)
        evento = resultado.manifestacoes[-1] if resultado.manifestacoes else None
        valores = [
            resultado.chave, int(numero) if numero else None, serie,
            DESCRICOES.get(evento.codigo, evento.descricao) if evento else "Sem evento localizado",
            _data_exibicao(evento.data) if evento else "",
            evento.protocolo if evento else "",
            evento.codigo if evento else "",
            len(resultado.manifestacoes),
            "\n".join(f"{_data_exibicao(e.data)} · {DESCRICOES.get(e.codigo, e.descricao)} · {e.protocolo}"
                      for e in resultado.manifestacoes),
            resultado.erro or "",
            status_label(resultado.cancelada),
            spreadsheet_label(resultado.ctes),
        ]
        for coluna, valor in enumerate(valores, 1):
            celula = ws.cell(indice, coluna, valor)
            if isinstance(valor, str):
                celula.data_type = "s"  # impede formula de conteudo externo
            style_cell(celula, alternate=indice % 2 == 0, wrap=coluna in (4, 9, 10), center=coluna in (3, 7, 8))
        for coluna in (1, 3, 6, 7):
            ws.cell(indice, coluna).number_format = "@"
        ws.cell(indice, 2).number_format = "0"
        ws.cell(indice, 8).number_format = "0"
        ws.row_dimensions[indice].height = max(text_row_height(valores[8], 78), text_row_height(valores[9], 40), text_row_height(valores[3], 34))

    configure_sheet(ws, header_row=4, widths=[51, 13, 10, 34, 31, 23, 13, 11, 78, 40, 19, 23], title=True)
    ws.auto_filter.ref = f"A4:L{max(4, 4 + len(resultados))}"

    if numeros_solicitados is not None:
        summary = wb.create_sheet("Numeros solicitados")
        summary.append(["Número NF", "Notas localizadas", "Resultado"])
        for number, count in cancellable_rows(numeros_solicitados.items()):
            summary.append([number, count, "Não localizada no banco" if not count else
                            "Localizada" if count == 1 else "Múltiplas notas; confira chave e série"])
        for row in cancellable_rows(summary.iter_rows(min_row=2)):
            for cell in row:
                style_cell(cell, alternate=cell.row % 2 == 0, wrap=cell.column == 3)
            summary.row_dimensions[row[0].row].height = 30
        configure_sheet(summary, header_row=1, widths=[18, 23, 54])
        summary.freeze_panes = "A2"
        summary.auto_filter.ref = summary.dimensions

    destino = Path(caminho)
    destino.parent.mkdir(parents=True, exist_ok=True)
    descritor, temporario = tempfile.mkstemp(prefix=".nfe-consulta-", suffix=".xlsx", dir=destino.parent)
    os.close(descritor)
    try:
        check_export_cancelled()
        wb.save(temporario)
        check_export_cancelled()
        os.replace(temporario, destino)
    finally:
        Path(temporario).unlink(missing_ok=True)
