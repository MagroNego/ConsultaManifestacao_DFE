"""Arquivo cumulativo de XMLs emitidos e relatórios do leitor, no banco fiscal."""
from __future__ import annotations

import csv
import hashlib
import json
import re
import stat
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP, localcontext
from io import BytesIO, StringIO
from pathlib import Path, PurePosixPath
from zipfile import BadZipFile, ZipFile

from defusedxml import ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from nfe_consulta.seguranca_banco import abrir_banco
from nfe_consulta.cte import linked_ctes, export_label
from nfe_consulta.cancelamento import SCHEMA as CANCELLATION_SCHEMA, read_cancellation, save_cancellation, cancelled_keys, status_label
from nfe_consulta.xml_reader import parse_nfe_xml, decimal_value, FIELDS_ITENS, FIELDS_RETIDO

MAX_XML = 8 * 1024 * 1024
MAX_UPLOAD = 100 * 1024 * 1024
MAX_EXPANDED = 250 * 1024 * 1024
MAX_FILES = 5000
MAX_EXPORT_ROWS = 100_000
FIELDS_NOTAS = ["Chave Acesso", "Numero Nota", "Serie", "Data Emissao", "Emitente Razao", "Destinatario Razao", "Vl. Nota", "Itens"]
EXCEL_EXCLUDED_COLUMNS = frozenset({
    "CSOSN", "Origem ICMS", "Qtd Base PIS", "Aliquota PIS por unidade",
    "Qtd Base COFINS", "Aliquota COFINS por unidade",
})


def _schema(conn):
    conn.executescript(CANCELLATION_SCHEMA)
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS xml_documentos (
            cnpj TEXT NOT NULL, chave TEXT NOT NULL, numero TEXT NOT NULL,
            serie TEXT NOT NULL, emissao TEXT NOT NULL, emitente TEXT NOT NULL,
            destinatario TEXT NOT NULL, valor TEXT NOT NULL, itens INTEGER NOT NULL,
            nome TEXT NOT NULL, sha256 TEXT NOT NULL, xml BLOB NOT NULL,
            importado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY(cnpj, chave)
        );
        CREATE INDEX IF NOT EXISTS xml_documentos_emissao ON xml_documentos(cnpj, emissao);
        CREATE TABLE IF NOT EXISTS xml_relatorio (
            cnpj TEXT NOT NULL, chave TEXT NOT NULL, tipo TEXT NOT NULL,
            linha INTEGER NOT NULL, dados TEXT NOT NULL, busca TEXT NOT NULL,
            PRIMARY KEY(cnpj, chave, tipo, linha),
            FOREIGN KEY(cnpj, chave) REFERENCES xml_documentos(cnpj, chave)
        );
        CREATE TABLE IF NOT EXISTS xml_importacoes (
            id INTEGER PRIMARY KEY, usuario TEXT NOT NULL,
            criado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, resultado TEXT NOT NULL
        );
    """)


def _has_table(conn, table):
    return conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)).fetchone() is not None


def _name(name):
    return PurePosixPath(str(name).replace("\\", "/")).name[:180] or "arquivo.xml"


def _text(parent, tag):
    if parent is None:
        return ""
    child = next((el for el in parent if el.tag.rsplit("}", 1)[-1] == tag), None)
    return (child.text or "").strip() if child is not None else ""


def _child(parent, tag):
    if parent is None:
        return None
    return next((el for el in parent if el.tag.rsplit("}", 1)[-1] == tag), None)


def read_document(data: bytes, cnpj: str):
    if len(data) > MAX_XML:
        raise ValueError("XML excede 8 MiB.")
    root = ET.fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
    if root.tag.rsplit("}", 1)[-1] not in ("NFe", "nfeProc"):
        raise ValueError("Envie NF-e completa; resumos e eventos não são XMLs da nota.")
    nfe = root if root.tag.rsplit("}", 1)[-1] == "NFe" else _child(root, "NFe")
    inf = _child(nfe, "infNFe")
    if inf is None:
        raise ValueError("XML sem infNFe.")
    chave = inf.attrib.get("Id", "").removeprefix("NFe")
    if not re.fullmatch(r"\d{44}", chave) or chave[20:22] != "55":
        raise ValueError("Chave NF-e modelo 55 inválida.")
    ide, emit, dest = (_child(inf, tag) for tag in ("ide", "emit", "dest"))
    if _text(emit, "CNPJ") != cnpj or chave[6:20] != cnpj:
        raise ValueError("O XML não foi emitido pelo CNPJ configurado no app.")
    numero, serie = _text(ide, "nNF"), _text(ide, "serie")
    if not numero.isdigit() or not serie.isdigit() or int(numero) != int(chave[25:34]) or int(serie) != int(chave[22:25]):
        raise ValueError("Número ou série diverge da chave da NF-e.")
    emissao = (_text(ide, "dhEmi") or _text(ide, "dEmi"))[:10]
    datetime.strptime(emissao, "%Y-%m-%d")
    rows, retained = parse_nfe_xml(BytesIO(data))
    if not rows or len(rows) > 990:
        raise ValueError("NF-e sem itens ou acima do limite de 990 itens.")
    valor = _text(_child(_child(inf, "total"), "ICMSTot"), "vNF")
    decimal_value(valor)
    meta = dict(chave=chave, numero=numero, serie=serie, emissao=emissao,
                emitente=_text(emit, "xNome"), destinatario=_text(dest, "xNome"), valor=valor)
    return meta, rows, retained


def _documents(paths):
    count = expanded = 0
    for path, filename in paths:
        suffix = Path(filename).suffix.lower()
        if suffix == ".xml":
            size = path.stat().st_size
            count += 1
            expanded += size
            if count > MAX_FILES or expanded > MAX_EXPANDED:
                raise ValueError("Lote excede 5.000 documentos ou 250 MiB descompactados.")
            if size > MAX_XML:
                yield _name(filename), None, "XML excede 8 MiB."
            else:
                yield _name(filename), path.read_bytes(), None
        elif suffix == ".zip":
            try:
                with ZipFile(path) as archive:
                    entries = archive.infolist()
                    if len(entries) > MAX_FILES:
                        raise ValueError("ZIP excede 5.000 entradas.")
                    for entry in entries:
                        if entry.is_dir():
                            continue
                        count += 1
                        expanded += entry.file_size
                        if count > MAX_FILES or expanded > MAX_EXPANDED:
                            raise ValueError("Lote excede 5.000 documentos ou 250 MiB descompactados.")
                        name = entry.filename.replace("\\", "/")
                        parts = PurePosixPath(name)
                        if parts.is_absolute() or ".." in parts.parts or ":" in name or stat.S_ISLNK(entry.external_attr >> 16):
                            yield _name(name), None, "Caminho inseguro no ZIP."
                            continue
                        if parts.suffix.lower() != ".xml":
                            yield _name(name), None, "Arquivo ignorado: envie XMLs de NF-e."
                            continue
                        if entry.flag_bits & 1 or entry.file_size > MAX_XML:
                            yield _name(name), None, "XML protegido por senha ou acima de 8 MiB."
                            continue
                        try:
                            with archive.open(entry) as stream:
                                data = stream.read(MAX_XML + 1)
                            if len(data) > MAX_XML:
                                yield _name(name), None, "XML excede 8 MiB."
                            else:
                                yield _name(name), data, None
                        except (BadZipFile, RuntimeError, NotImplementedError, OSError):
                            yield _name(name), None, "Não foi possível ler este XML no ZIP."
            except BadZipFile as exc:
                raise ValueError("Arquivo ZIP inválido.") from exc
        else:
            yield _name(filename), None, "Arquivo ignorado: envie XML ou ZIP."


def import_batch(database_path, cnpj, paths, *, password=None, user="admin"):
    if not Path(database_path).is_file():
        raise ValueError("Configure o banco existente antes de importar.")
    conn = abrir_banco(database_path, password)
    summary = dict(importadas=0, duplicadas=0, cancelamentos=0, erros=0, itens=0, retencoes=0, registros=[])
    try:
        _schema(conn)
        # Um erro estrutural de ZIP/limite desfaz o lote; erros individuais ficam no registro.
        with conn:
            for name, data, error in _documents(paths):
                if error:
                    summary["erros"] += 1
                    summary["registros"].append(dict(arquivo=name, status="Ignorado", mensagem=error))
                    continue
                try:
                    root = ET.fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
                    cancellation = read_cancellation(root, cnpj)
                    event_only = cancellation is not None and root.tag.rsplit('}', 1)[-1] != 'nfeProc'
                    if not event_only:
                        meta, rows, retained = read_document(data, cnpj)
                except Exception:
                    summary["erros"] += 1
                    summary["registros"].append(dict(arquivo=name, status="Erro", mensagem="XML inválido, incompleto ou de outro emitente."))
                    continue
                if event_only:
                    added = save_cancellation(conn, cnpj, cancellation, data)
                    summary['cancelamentos' if added else 'duplicadas'] += 1
                    summary['registros'].append(dict(arquivo=name, status='Cancelamento' if added else 'Duplicada',
                        mensagem='Cancelamento registrado pela chave da NF-e.' if added else 'Cancelamento já registrado.'))
                    continue
                chave = meta["chave"]
                if cancellation is not None:
                    summary['cancelamentos'] += save_cancellation(conn, cnpj, cancellation, data)
                conn.execute('''INSERT INTO informacoes_nfe(cnpj,chave,emitente) VALUES (?,?,?)
                    ON CONFLICT(cnpj,chave) DO UPDATE SET emitente=CASE WHEN informacoes_nfe.emitente='' THEN excluded.emitente ELSE informacoes_nfe.emitente END''', (cnpj, chave, meta['emitente']))
                if conn.execute("SELECT 1 FROM xml_documentos WHERE cnpj=? AND chave=?", (cnpj, chave)).fetchone():
                    summary["duplicadas"] += 1
                    summary["registros"].append(dict(arquivo=name, status="Duplicada", mensagem="Chave já arquivada; original preservado."))
                    continue
                conn.execute("""INSERT INTO xml_documentos
                    (cnpj,chave,numero,serie,emissao,emitente,destinatario,valor,itens,nome,sha256,xml)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""", (cnpj, chave, meta["numero"], meta["serie"], meta["emissao"], meta["emitente"], meta["destinatario"], meta["valor"], len(rows), name, hashlib.sha256(data).hexdigest(), data))
                for kind, report in (("itens", rows), ("retencoes", [retained] if retained else [])):
                    for index, row in enumerate(report):
                        row["Chave Acesso"] = chave
                        conn.execute("INSERT INTO xml_relatorio VALUES (?,?,?,?,?,?)", (cnpj, chave, kind, index, json.dumps(row, ensure_ascii=False), " ".join(str(v) for v in row.values()).casefold()))
                summary["importadas"] += 1
                summary["itens"] += len(rows)
                summary["retencoes"] += bool(retained)
                summary["registros"].append(dict(arquivo=name, status="Importada", mensagem=f"NF {meta['numero']} · {len(rows)} item(ns)."))
            cur = conn.execute("INSERT INTO xml_importacoes(usuario,resultado) VALUES (?,?)", (user, json.dumps(summary, ensure_ascii=False)))
            summary["id"] = cur.lastrowid
        return summary
    finally:
        conn.close()


def existing_keys(database_path, cnpj, keys, *, password=None):
    if not keys or not Path(database_path).is_file():
        return set()
    conn = abrir_banco(database_path, password, somente_leitura=True)
    try:
        if not _has_table(conn, "xml_documentos"):
            return set()
        marks = ",".join("?" for _ in keys)
        return {row[0] for row in conn.execute(f"SELECT chave FROM xml_documentos WHERE cnpj=? AND chave IN ({marks})", (cnpj, *keys))}
    finally:
        conn.close()


def download_xml(database_path, cnpj, chave, *, password=None):
    conn = abrir_banco(database_path, password, somente_leitura=True)
    try:
        if not _has_table(conn, "xml_documentos"):
            return None
        row = conn.execute("SELECT xml FROM xml_documentos WHERE cnpj=? AND chave=?", (cnpj, chave)).fetchone()
        return bytes(row[0]) if row else None
    finally:
        conn.close()


def query_report(database_path, cnpj, *, password=None, kind="notas", query="", inicio=None, fim=None, page=1, export=False):
    if kind not in ("notas", "itens", "retencoes") or page < 1 or len(query) > 200:
        raise ValueError("Filtro do leitor inválido.")
    fields = FIELDS_ITENS if kind == "itens" else FIELDS_RETIDO if kind == "retencoes" else FIELDS_NOTAS
    columns = [*fields, "Situação", "CT-e vinculado"]
    empty = dict(rows=[], columns=columns, total=0, page=1, pages=0, notas=0)
    if not Path(database_path).is_file():
        return empty
    conn = abrir_banco(database_path, password, somente_leitura=True)
    try:
        if not _has_table(conn, "xml_documentos"):
            return empty
        params = [cnpj]
        clauses = ["d.cnpj=?"]
        for field, value in (("d.emissao>=?", inicio), ("d.emissao<=?", fim)):
            if value:
                clauses.append(field)
                params.append(value.isoformat())
        source = "xml_documentos d"
        if kind != "notas":
            source += " JOIN xml_relatorio r ON r.cnpj=d.cnpj AND r.chave=d.chave"
            clauses.append("r.tipo=?")
            params.append(kind)
        if query.strip():
            if kind == "notas":
                clauses.append("(d.chave LIKE ? OR d.numero LIKE ? OR LOWER(d.destinatario) LIKE ? OR LOWER(d.emitente) LIKE ?)")
                params.extend(["%" + query.strip().casefold() + "%"] * 4)
            else:
                clauses.append("r.busca LIKE ?")
                params.append("%" + query.strip().casefold() + "%")
        where = " AND ".join(clauses)
        total = conn.execute(f"SELECT COUNT(*) FROM {source} WHERE {where}", params).fetchone()[0]
        if export and total > MAX_EXPORT_ROWS:
            raise ValueError("Relatório excede 100.000 linhas. Refine os filtros.")
        pages = (total + 99) // 100
        page = min(page, max(1, pages))
        select = "d.chave,d.numero,d.serie,d.emissao,d.emitente,d.destinatario,d.valor,d.itens" if kind == "notas" else "r.dados"
        suffix = " ORDER BY d.emissao DESC,d.chave" + (",r.linha" if kind != "notas" else "")
        if not export:
            suffix += " LIMIT 100 OFFSET ?"
        lines = conn.execute(f"SELECT {select} FROM {source} WHERE {where}" + suffix, params if export else [*params, (page - 1) * 100]).fetchall()
        rows = [dict(zip(FIELDS_NOTAS, row)) for row in lines] if kind == "notas" else [json.loads(row[0]) for row in lines]
        ctes = linked_ctes(conn, cnpj, [row.get("Chave Acesso", "") for row in rows])
        cancelled = cancelled_keys(conn, cnpj, [row.get('Chave Acesso', '') for row in rows])
        for row in rows:
            row['Situação'] = status_label(row.get('Chave Acesso', '') in cancelled)
            row["CT-e vinculado"] = export_label(ctes.get(row.get("Chave Acesso", ""), ()))
        return dict(rows=rows, columns=columns, total=total, page=page, pages=pages, notas=conn.execute("SELECT COUNT(*) FROM xml_documentos WHERE cnpj=?", (cnpj,)).fetchone()[0])
    finally:
        conn.close()


def import_history(database_path, *, password=None, batch_id=None):
    if not Path(database_path).is_file():
        return [], None
    conn = abrir_banco(database_path, password, somente_leitura=True)
    try:
        if not _has_table(conn, "xml_importacoes"):
            return [], None
        history = [dict(id=row[0], usuario=row[1], criado_em=row[2], **{k: v for k, v in json.loads(row[3]).items() if k != "registros"}) for row in conn.execute("SELECT id,usuario,criado_em,resultado FROM xml_importacoes ORDER BY id DESC LIMIT 20")]
        row = conn.execute("SELECT resultado FROM xml_importacoes WHERE id=?", (batch_id,)).fetchone() if batch_id else None
        return history, json.loads(row[0]) if row else None
    finally:
        conn.close()


def export_csv(report):
    out = StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=report["columns"], delimiter=";", quoting=csv.QUOTE_ALL)
    writer.writeheader()
    for row in report["rows"]:
        writer.writerow({key: ("'" + str(row.get(key, "")) if str(row.get(key, "")).lstrip().startswith(("=", "+", "-", "@")) else row.get(key, "")) for key in report["columns"]})
    return out.getvalue().encode("utf-8-sig")


def export_excel(reports):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, report in reports:
        sheet = workbook.create_sheet(title)
        columns = [key for key in report["columns"] if key not in EXCEL_EXCLUDED_COLUMNS]
        sheet.append(columns)
        for row in report["rows"]:
            values = []
            for key in columns:
                value = row.get(key, "")
                if key == "Data Emissao" and value:
                    for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
                        try:
                            value = datetime.strptime(str(value), pattern).date()
                            break
                        except ValueError:
                            continue
                elif key.startswith(("Vl.", "Base ", "% ", "Qtd", "Aliquota ")) and value != "":
                    try:
                        with localcontext() as context:
                            context.prec = 110
                            value = Decimal(str(value).replace(",", ".")).quantize(
                                Decimal("0.0001"), rounding=ROUND_HALF_UP)
                    except Exception:
                        pass
                values.append(value)
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, date):
                    cell.number_format = "dd/mm/yyyy"
                elif isinstance(cell.value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                elif cell.value is not None:
                    cell.number_format = "#,##0" if cell.value == int(cell.value) else "#,##0.0000"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="193D65")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, key in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 48 if "Chave" in key else 36 if "Razao" in key or key == "Descricao" else 18
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()
