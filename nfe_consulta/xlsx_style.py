"""Formatação compartilhada dos relatórios Excel, sem alterar seus dados."""
from math import ceil

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BODY_FONT = Font(name='Calibri', size=11, color='20252B')
HEADER_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='343A40')
HEADER_BORDER = Border(bottom=Side(style='thin', color='AEB4BA'))
ROW_FILL = PatternFill('solid', fgColor='F3F4F5')
BODY_ALIGNMENT = Alignment(vertical='center')
WRAPPED_ALIGNMENT = Alignment(vertical='center', wrap_text=True)
CENTER_ALIGNMENT = Alignment(horizontal='center', vertical='center')


def style_cell(cell, *, alternate=False, wrap=False, center=False):
    cell.font = BODY_FONT
    cell.alignment = CENTER_ALIGNMENT if center else WRAPPED_ALIGNMENT if wrap else BODY_ALIGNMENT
    if alternate:
        cell.fill = ROW_FILL


def configure_sheet(sheet, *, header_row, widths, title=False):
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    for cell in sheet[header_row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[header_row].height = 30
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    if title:
        last = get_column_letter(len(widths))
        sheet.merge_cells(f'A1:{last}1')
        sheet.merge_cells(f'A2:{last}2')
        sheet['A1'].font = Font(name='Calibri', size=16, bold=True, color='20252B')
        sheet['A1'].alignment = BODY_ALIGNMENT
        sheet.row_dimensions[1].height = 32
        sheet['A2'].font = Font(name='Calibri', size=11, color='59616A')
        sheet['A2'].alignment = WRAPPED_ALIGNMENT
        sheet.row_dimensions[2].height = 30
        sheet.row_dimensions[3].height = 9
    sheet.print_title_rows = f'1:{header_row}'
    sheet.print_options.horizontalCentered = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def text_row_height(text, width, *, minimum=30):
    """Reserva espaço para quebras explícitas e linhas longas, até o limite do Excel."""
    lines = sum(max(1, ceil(len(line) / max(1, width - 3))) for line in str(text).split('\n'))
    return min(409, max(minimum, 17 * lines + 10))
