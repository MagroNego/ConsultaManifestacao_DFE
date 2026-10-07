from pathlib import Path
from html import unescape
from io import BytesIO
import re
import runpy
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from nfe_consulta.web.app import create_app

h = runpy.run_path(str(Path(__file__).with_name('test_filtro_situacao.py')))


def test_manifestations_detail_is_scoped_and_back_keeps_emission_filters(tmp_path):
    cfg = h['setup'](tmp_path)
    key, other = h['KEY'], h['KEY2']
    with TestClient(create_app(cfg)) as client:
        original = '/xml?q=123&pagina=2&situacao=cancelada'
        page = client.get(original).text
        href = unescape(re.search(r'data-items-open href="([^"]+)" title="Manifestações"', page).group(1))
        response = client.get(href)
        assert response.status_code == 200
        assert '<h1>MANIFESTAÇÕES | NF 123</h1>' in response.text
        assert 'href="' + original.replace('&', '&amp;') + '" data-items-back' in response.text
        assert 'Consulta por arquivo' not in response.text
        rows = response.text.split('<tbody>')[1].split('</tbody>')[0]
        assert key in rows and other not in rows
        assert 'name="chave" value="' + key + '"' in response.text
        filtered = client.get('/xml/manifestacoes/' + key, params={'codigo': '210200', 'voltar': original})
        assert 'Nenhum evento encontrado' in filtered.text
        assert 'name="voltar" value="' in filtered.text
        cleared = re.search(r'href="([^"]+)"\>Limpar', filtered.text).group(1)
        assert 'voltar=' in cleared
        exported = client.get('/consulta/exportar', params={'chave': key})
        wb = load_workbook(BytesIO(exported.content))
        assert wb['Manifestacoes'].max_row == 5
        assert client.get('/xml/manifestacoes/invalid').status_code == 400
        assert client.get('/xml/manifestacoes/' + '0' * 44).status_code == 404
        assert client.get('/xml/manifestacoes/' + key + '?voltar=https://evil.example').text.count('href="/xml" data-items-back') == 1


def test_imported_note_without_manifestations_has_dedicated_empty_screen(tmp_path):
    items = runpy.run_path(str(Path(__file__).with_name('test_xml_items.py')))
    cfg = items['setup'](tmp_path)
    with TestClient(create_app(cfg)) as client:
        response = client.get('/xml/manifestacoes/' + items['KEY'])
        assert response.status_code == 200
        assert 'MANIFESTAÇÕES | NF 123' in response.text
        assert 'Nenhum evento encontrado' in response.text
