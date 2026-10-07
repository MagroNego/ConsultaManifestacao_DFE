from io import BytesIO
from pathlib import Path
import runpy
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from nfe_consulta.banco import BancoManifestacoes
from nfe_consulta.config import CNPJ_PADRAO
from nfe_consulta.modelos import Manifestacao, RetornoDistribuicao, InformacaoNota
from nfe_consulta.web.app import create_app
from nfe_consulta.web import xml_store
from nfe_consulta.web.consulta_local import consultar_eventos, normalizar_filtros
from nfe_consulta.seguranca_banco import abrir_banco

helpers = runpy.run_path(str(Path(__file__).with_name('test_xml_integrado.py')))
web = helpers['helpers']
KEY = helpers['CHAVE']


def cancellation(status='135'):
    return f'''<procEventoNFe xmlns="http://www.portalfiscal.inf.br/nfe" versao="1.00">
    <evento versao="1.00"><infEvento><tpAmb>1</tpAmb><CNPJ>{CNPJ_PADRAO}</CNPJ><chNFe>{KEY}</chNFe>
    <dhEvento>2026-10-07T09:00:00-03:00</dhEvento><tpEvento>110111</tpEvento><nSeqEvento>1</nSeqEvento>
    <detEvento><descEvento>Cancelamento</descEvento><nProt>135260000000001</nProt></detEvento></infEvento></evento>
    <retEvento versao="1.00"><infEvento><tpAmb>1</tpAmb><cStat>{status}</cStat><chNFe>{KEY}</chNFe>
    <tpEvento>110111</tpEvento><nSeqEvento>1</nSeqEvento><dhRegEvento>2026-10-07T09:01:00-03:00</dhRegEvento>
    <nProt>135260000000002</nProt></infEvento></retEvento></procEventoNFe>'''.encode()


def batch(cfg, tmp_path, payloads):
    paths = []
    for index, payload in enumerate(payloads):
        path = tmp_path / f'input{index}.xml'
        path.write_bytes(payload)
        paths.append((path, path.name))
    return xml_store.import_batch(cfg.database_path, CNPJ_PADRAO, paths)


@pytest.mark.parametrize('event_first', [True, False])
@pytest.mark.parametrize('status', ['135', '155'])
def test_cancellation_order_idempotence_and_no_reactivation(tmp_path, event_first, status):
    cfg = web['settings_web'](tmp_path)
    web['criar_banco'](cfg.database_path)
    payloads = [cancellation(status), helpers['xml']()]
    if not event_first:
        payloads.reverse()
    result = batch(cfg, tmp_path, payloads)
    assert result['cancelamentos'] == result['importadas'] == 1 and result['erros'] == 0
    assert batch(cfg, tmp_path, payloads)['duplicadas'] == 2
    report = xml_store.query_report(cfg.database_path, CNPJ_PADRAO)
    assert report['rows'][0]['Situação'] == 'Cancelada'
    assert xml_store.download_xml(cfg.database_path, CNPJ_PADRAO, KEY) == helpers['xml']()
    banco = BancoManifestacoes(str(cfg.database_path))
    banco.salvar_retorno(CNPJ_PADRAO, RetornoDistribuicao(138, 'ok', '2', '2', (),
        informacoes_notas=(InformacaoNota(KEY, 'Emitente', False),)))
    assert banco.consultar_chave(KEY, CNPJ_PADRAO).cancelada
    assert banco.consultar_chave(KEY).cancelada
    assert not banco.consultar_chave(KEY, '00000000000000').cancelada
    banco.fechar()


@pytest.mark.parametrize('payload', [
    cancellation('136'), cancellation('204'),
    cancellation().replace(b'<tpAmb>1</tpAmb>', b'<tpAmb>2</tpAmb>'),
    cancellation().replace(b'<CNPJ>' + CNPJ_PADRAO.encode(), b'<CNPJ>00000000000000'),
    cancellation().replace(b'<tpEvento>110111</tpEvento>', b'<tpEvento>110110</tpEvento>'),
    cancellation().replace(b'<nSeqEvento>1</nSeqEvento>', b'<nSeqEvento>2</nSeqEvento>', 1),
    cancellation().replace(KEY.encode(), (KEY[:-1] + ('1' if KEY[-1] != '1' else '2')).encode(), 1),
    cancellation().replace(b'135260000000002', b'123'),
    cancellation().replace(b'<retEvento', b'<semRetorno').replace(b'</retEvento>', b'</semRetorno>'),
    b'<!DOCTYPE procEventoNFe>' + cancellation(),
])
def test_request_rejection_invalid_protocol_and_mismatch_do_not_cancel(tmp_path, payload):
    cfg = web['settings_web'](tmp_path)
    web['criar_banco'](cfg.database_path)
    result = batch(cfg, tmp_path, [helpers['xml'](), payload])
    assert result['importadas'] == 1 and result['erros'] == 1 and result['cancelamentos'] == 0
    assert xml_store.query_report(cfg.database_path, CNPJ_PADRAO)['rows'][0]['Situação'] == 'Sem cancelamento registrado'


def test_all_ui_and_reports_use_same_cancellation_flag(tmp_path):
    cfg = web['settings_web'](tmp_path)
    banco = BancoManifestacoes(str(cfg.database_path))
    banco.salvar_manifestacoes(CNPJ_PADRAO, RetornoDistribuicao(138, 'ok', '1', '1',
        ((KEY, Manifestacao('210240', 'Operação não Realizada', '2026-10-01T10:00:00-03:00', '135260000000001')),)))
    banco.fechar()
    batch(cfg, tmp_path, [helpers['xml'](), cancellation()])
    events = consultar_eventos(cfg.database_path, CNPJ_PADRAO, normalizar_filtros(chave=KEY))
    assert events.eventos[0].cancelada and events.eventos[0].descricao == 'Operação não Realizada'
    app = create_app(cfg)
    with TestClient(app) as client:
        for url in ['/consulta?chave=' + KEY, '/xml', '/xml?tipo=itens', '/xml?tipo=retencoes']:
            response = client.get(url)
            assert response.status_code == 200
            assert 'Cancelada' in response.text and '>Emissão</a>' in response.text
        for kind in ['notas', 'itens', 'retencoes']:
            report = client.get('/xml/exportar?formato=csv&tipo=' + kind).content.decode('utf-8-sig')
            assert 'Situação' in report and 'Cancelada' in report
        for url, sheet, header_row in [('/xml/exportar?formato=xlsx', 'Notas', 1), ('/consulta/exportar?chave=' + KEY, 'Manifestacoes', 4)]:
            response = client.get(url)
            ws = load_workbook(BytesIO(response.content))[sheet]
            headers = [cell.value for cell in ws[header_row]]
            assert ws.cell(header_row + 1, headers.index('Situação') + 1).value == 'Cancelada'
        # Importar um evento não altera a sequência NSU nem inventa manifestação.
        conn = abrir_banco(cfg.database_path, None, somente_leitura=True)
        assert conn.execute('SELECT COUNT(*) FROM manifestacoes').fetchone()[0] == 1
        conn.close()


def test_cancellation_without_note_or_manifestation_is_visible_on_lookup(tmp_path):
    cfg = web['settings_web'](tmp_path)
    web['criar_banco'](cfg.database_path)
    assert batch(cfg, tmp_path, [cancellation()])['cancelamentos'] == 1
    with TestClient(create_app(cfg)) as client:
        for query in ['chave=' + KEY, 'numero=' + str(int(KEY[25:34]))]:
            page = client.get('/consulta?' + query)
            assert 'Cancelada' in page.text and KEY in page.text
            assert 'Nenhum evento encontrado' in page.text


def test_legacy_cancelled_protocol(tmp_path):
    cfg = web['settings_web'](tmp_path)
    web['criar_banco'](cfg.database_path)
    for status in ('101', '151'):
        content = helpers['xml']().replace(b'</nfeProc>', f'<protNFe><infProt><tpAmb>1</tpAmb><chNFe>{KEY}</chNFe><cStat>{status}</cStat><nProt>135260000000003</nProt><dhRecbto>2026-10-07T09:01:00-03:00</dhRecbto></infProt></protNFe></nfeProc>'.encode())
        assert batch(cfg, tmp_path, [content])['erros'] == 0
    assert xml_store.query_report(cfg.database_path, CNPJ_PADRAO)['rows'][0]['Situação'] == 'Cancelada'


def test_zip_and_structural_failure_roll_back_cancellation(tmp_path):
    cfg = web['settings_web'](tmp_path)
    web['criar_banco'](cfg.database_path)
    zip_path = tmp_path / 'lote.zip'
    with ZipFile(zip_path, 'w') as archive:
        archive.writestr('cancelamento.xml', cancellation())
        archive.writestr('nota.xml', helpers['xml']())
    result = xml_store.import_batch(cfg.database_path, CNPJ_PADRAO, [(zip_path, 'lote.zip')])
    assert result['cancelamentos'] == result['importadas'] == 1
    cfg2 = web['settings_web'](tmp_path / 'rollback')
    web['criar_banco'](cfg2.database_path)
    event = tmp_path / 'event.xml'
    event.write_bytes(cancellation())
    broken = tmp_path / 'broken.zip'
    broken.write_bytes(b'invalid zip')
    with pytest.raises(ValueError):
        xml_store.import_batch(cfg2.database_path, CNPJ_PADRAO, [(event, 'event.xml'), (broken, 'broken.zip')])
    banco = BancoManifestacoes(str(cfg2.database_path))
    assert not banco.consultar_chave(KEY, CNPJ_PADRAO).cancelada
    banco.fechar()


def test_cancellation_in_encrypted_database(tmp_path):
    path = tmp_path / 'secure.db'
    password = 'test-only-encrypted-password'
    BancoManifestacoes(str(path), password).fechar()
    files = []
    for index, data in enumerate([helpers['xml'](), cancellation()]):
        file = tmp_path / f'encrypted{index}.xml'
        file.write_bytes(data)
        files.append((file, file.name))
    result = xml_store.import_batch(path, CNPJ_PADRAO, files, password=password)
    assert result['cancelamentos'] == 1
    assert xml_store.query_report(path, CNPJ_PADRAO, password=password)['rows'][0]['Situação'] == 'Cancelada'


def test_existing_nsu_flag_applies_to_new_xml(tmp_path):
    cfg = web['settings_web'](tmp_path)
    # Um cancelamento já recebido via NSU também vale para XML importado depois.
    banco = BancoManifestacoes(str(cfg.database_path))
    banco.salvar_retorno(CNPJ_PADRAO, RetornoDistribuicao(138, 'ok', '2', '2', (),
        informacoes_notas=(InformacaoNota(KEY, 'Emitente', True),)))
    banco.fechar()
    batch(cfg, tmp_path, [helpers['xml']()])
    assert xml_store.query_report(cfg.database_path, CNPJ_PADRAO)['rows'][0]['Situação'] == 'Cancelada'


@pytest.mark.parametrize('container', ['nfeProc', 'retConsSitNFe'])
def test_cancellation_embedded_in_note_or_status_response(tmp_path, container):
    cfg = web['settings_web'](tmp_path)
    web['criar_banco'](cfg.database_path)
    if container == 'nfeProc':
        data = helpers['xml']().replace(b'</nfeProc>', cancellation() + b'</nfeProc>')
    else:
        data = f'<retConsSitNFe xmlns="http://www.portalfiscal.inf.br/nfe"><cStat>101</cStat><chNFe>{KEY}</chNFe>'.encode() + cancellation() + b'</retConsSitNFe>'
    result = batch(cfg, tmp_path, [data])
    assert result['cancelamentos'] == 1 and result['erros'] == 0
    banco = BancoManifestacoes(str(cfg.database_path))
    assert banco.consultar_chave(KEY, CNPJ_PADRAO).cancelada
    banco.fechar()
