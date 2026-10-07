import base64
import gzip
from io import BytesIO
from pathlib import Path
import runpy
import sqlite3

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from nfe_consulta.banco import BancoManifestacoes
from nfe_consulta.config import CNPJ_PADRAO
from nfe_consulta.cte import linked_ctes
from nfe_consulta.modelos import Manifestacao, RetornoDistribuicao, NfeErroResposta
from nfe_consulta.parser_distribuicao import parse_retorno_distribuicao
from nfe_consulta.web.app import create_app
from nfe_consulta.web.consulta_local import consultar_eventos, FiltrosEventos
from nfe_consulta.web import xml_store

helpers = runpy.run_path(str(Path(__file__).with_name('test_xml_integrado.py')))
web = helpers['helpers']
NFE = helpers['CHAVE']
CTE = NFE[:20] + '57' + NFE[22:]
CTE2 = CTE[:25] + '000000456' + CTE[34:]


def event(code='610600', key=CTE, *, summary=False, status='135', protocol='891000000000001', sequence='1', date='2026-10-01T10:00:00-03:00'):
    body = f'<chNFe>{NFE}</chNFe><tpEvento>{code}</tpEvento><dhEvento>{date}</dhEvento><nSeqEvento>{sequence}</nSeqEvento>'
    if summary:
        return f'<resEvento xmlns="http://www.portalfiscal.inf.br/nfe">{body}<nProt>{protocol}</nProt></resEvento>'
    return f'''<procEventoNFe xmlns="http://www.portalfiscal.inf.br/nfe"><evento><infEvento>{body}
      <detEvento><CTe><chCTe>{key}</chCTe><nProt>333000000000001</nProt></CTe><emit><CNPJ>12345678000199</CNPJ><xNome>Transportadora teste</xNome></emit></detEvento>
      </infEvento></evento><retEvento><infEvento><cStat>{status}</cStat><chNFe>{NFE}</chNFe><tpEvento>{code}</tpEvento><nProt>{protocol}</nProt></infEvento></retEvento></procEventoNFe>'''


def distributed(*events):
    entries = ''.join(f'<docZip NSU="{i}" schema="resEvento_v1.01.xsd">{base64.b64encode(gzip.compress(xml.encode())).decode()}</docZip>' for i, xml in enumerate(events, 1))
    return parse_retorno_distribuicao(f'<retDistDFeInt><cStat>138</cStat><ultNSU>{len(events)}</ultNSU><maxNSU>{len(events)}</maxNSU><loteDistDFeInt>{entries}</loteDistDFeInt></retDistDFeInt>')


def test_full_and_summary_events_keep_separate_manifestations_and_event_protocol():
    result = distributed(event(), event(summary=True, sequence='2'))
    assert len(result.eventos_cte) == 2
    assert result.manifestacoes == ()
    assert result.documentos_ignorados == 0
    assert result.eventos_cte[0].protocolo == '891000000000001'
    assert result.eventos_cte[0].chave_cte == CTE
    assert result.eventos_cte[0].cnpj_transportadora == '12345678000199'
    assert result.eventos_cte[1].chave_cte == ''


@pytest.mark.parametrize('key', ['123', NFE, 'A' * 44])
def test_invalid_cte_key_rejects_batch(key):
    with pytest.raises(NfeErroResposta, match='CT-e inválida'):
        distributed(event(key=key))


def test_rejected_cte_event_not_recorded():
    result = distributed(event(status='573'))
    assert not result.eventos_cte
    assert result.documentos_ignorados == 1


@pytest.mark.parametrize('save_method', ['salvar_retorno', 'salvar_manifestacoes'])
def test_multiple_links_cancel_only_matching_cte_and_deduplicate(tmp_path, save_method):
    banco = BancoManifestacoes(str(tmp_path / 'db.sqlite'))
    result = distributed(event('610601', protocol='891000000000003'), event(), event(key=CTE2, sequence='2', protocol='891000000000002'))
    save = getattr(banco, save_method)
    assert save(CNPJ_PADRAO, result) == 0
    assert save(CNPJ_PADRAO, result) == 0
    assert banco.conexao.execute('SELECT COUNT(*) FROM eventos_cte').fetchone()[0] == 3
    links = banco.consultar_chave(NFE, CNPJ_PADRAO).ctes
    assert {link.chave: link.status for link in links} == {CTE: 'Cancelado', CTE2: 'Autorizado'}
    assert linked_ctes(banco.conexao, 'outro-cnpj', [NFE])[NFE] == ()
    banco.fechar()


def test_summary_cancellation_does_not_cancel_known_cte_without_key(tmp_path):
    banco = BancoManifestacoes(str(tmp_path / 'db.sqlite'))
    banco.salvar_manifestacoes(CNPJ_PADRAO, distributed(event(), event('610601', summary=True, protocol='891000000000003')))
    links = banco.consultar_chave(NFE, CNPJ_PADRAO).ctes
    assert next(link for link in links if link.chave == CTE).status == 'Autorizado'
    assert next(link for link in links if not link.chave).status == 'Cancelado'
    banco.fechar()


def test_summary_enriched_by_full_event_without_duplicate_link(tmp_path):
    banco = BancoManifestacoes(str(tmp_path / 'db.sqlite'))
    banco.salvar_manifestacoes(CNPJ_PADRAO, distributed(event(summary=True), event()))
    assert len(banco.consultar_chave(NFE, CNPJ_PADRAO).ctes) == 1
    banco.fechar()


def test_old_database_can_be_read_without_migration(tmp_path):
    cfg = web['settings_web'](tmp_path)
    web['criar_banco'](cfg.database_path)
    with sqlite3.connect(cfg.database_path) as conn:
        conn.execute('DROP TABLE eventos_cte')
        assert linked_ctes(conn, CNPJ_PADRAO, [NFE])[NFE] == []
    assert consultar_eventos(cfg.database_path, CNPJ_PADRAO, FiltrosEventos()).total == 0
    banco = BancoManifestacoes(str(cfg.database_path))
    assert linked_ctes(banco.conexao, CNPJ_PADRAO, [NFE])[NFE] == ()
    banco.fechar()


def test_links_in_public_ui_excel_csv_and_txt_lookup(tmp_path):
    cfg = web['settings_web'](tmp_path)
    banco = BancoManifestacoes(str(cfg.database_path))
    banco.salvar_retorno(CNPJ_PADRAO, distributed(event(), event('610601', protocol='891000000000003')))
    banco.salvar_manifestacoes(CNPJ_PADRAO, RetornoDistribuicao(138, 'ok', '2', '2', ((NFE, Manifestacao('210240', 'Operacao nao Realizada', '2026-10-02T10:00:00-03:00', '123')),)))
    banco.fechar()
    source = tmp_path / 'nota.xml'
    source.write_bytes(helpers['xml']())
    xml_store.import_batch(cfg.database_path, CNPJ_PADRAO, [(source, 'nota.xml')])
    app = create_app(cfg)
    with TestClient(app) as client:
        for url in ['/consulta?chave=' + NFE, '/xml']:
            response = client.get(url)
            assert response.status_code == 200
            assert CTE in response.text and 'CT-e 123 · Cancelado' in response.text
            assert 'Transportadora teste' in response.text
        csv = client.get('/xml/exportar?formato=csv&tipo=notas').content.decode('utf-8-sig')
        assert CTE in csv and 'Cancelado' in csv
        for url, sheet, header_row in [('/xml/exportar?formato=xlsx', 'Notas', 1), ('/consulta/exportar', 'Manifestacoes', 4)]:
            response = client.get(url)
            assert response.status_code == 200
            ws = load_workbook(BytesIO(response.content))[sheet]
            headers = [cell.value for cell in ws[header_row]]
            value = ws.cell(header_row + 1, headers.index('CT-e vinculado') + 1).value
            assert CTE in value and 'Cancelado' in value
    banco = BancoManifestacoes(str(cfg.database_path))
    assert len(banco.consultar_chave(NFE, CNPJ_PADRAO).ctes) == 1
    banco.fechar()


def test_invalid_cursor_does_not_save_cte_or_change_state(tmp_path):
    banco = BancoManifestacoes(str(tmp_path / 'db.sqlite'))
    banco.salvar_retorno(CNPJ_PADRAO, RetornoDistribuicao(137, 'ok', '000000000000010', '000000000000010', ()))
    with pytest.raises(NfeErroResposta, match='Cursor'):
        banco.salvar_retorno(CNPJ_PADRAO, distributed(event()))
    assert not banco.consultar_chave(NFE, CNPJ_PADRAO).ctes
    assert banco.obter_estado(CNPJ_PADRAO)[0] == '000000000000010'
    banco.fechar()


def test_txt_workbook_includes_cte_preserves_manifestation(tmp_path):
    from nfe_consulta.xlsx_writer import gravar_xlsx
    banco = BancoManifestacoes(str(tmp_path / 'db.sqlite'))
    banco.salvar_manifestacoes(CNPJ_PADRAO, distributed(event(summary=True)))
    result = banco.consultar_chave(NFE, CNPJ_PADRAO)
    banco.fechar()
    output = tmp_path / 'saida.xlsx'
    gravar_xlsx(str(output), [result], 'local')
    ws = load_workbook(output).active
    assert ws['D5'].value == 'Sem evento localizado'
    assert ws['K4'].value == 'CT-e vinculado'
    assert ws['K5'].value == 'Vínculo identificado · chave não informada'
    assert ws['K5'].data_type == 's'


def test_xml_empty_installation_keeps_empty_state(tmp_path):
    cfg = web['settings_web'](tmp_path)
    with TestClient(create_app(cfg)) as client:
        response = client.get('/xml')
        assert response.status_code == 200
        assert 'Nenhum XML corresponde aos filtros' in response.text
        assert 'Não foi possível abrir' not in response.text


def test_cte_saved_in_encrypted_database_without_plaintext_copy(tmp_path):
    from nfe_consulta.seguranca_banco import migrar_banco
    original = tmp_path / 'original.db'
    secure = tmp_path / 'seguro.db'
    banco = BancoManifestacoes(str(original))
    banco.fechar()
    password = 'Senha-Teste-CTe-123!'
    migrar_banco(original, secure, password)
    banco = BancoManifestacoes(str(secure), password)
    banco.salvar_manifestacoes(CNPJ_PADRAO, distributed(event()))
    assert banco.consultar_chave(NFE, CNPJ_PADRAO).ctes[0].chave == CTE
    banco.fechar()
    raw = secure.read_bytes()
    assert CTE.encode() not in raw and not raw.startswith(b'SQLite format 3')
    banco = BancoManifestacoes(str(original))
    assert not banco.consultar_chave(NFE, CNPJ_PADRAO).ctes
    banco.fechar()
