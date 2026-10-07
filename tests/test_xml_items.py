import csv
from html import unescape
from io import BytesIO, StringIO
from pathlib import Path
import re
import runpy

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from nfe_consulta.config import CNPJ_PADRAO
from nfe_consulta.web.app import create_app
from nfe_consulta.web import xml_store

h = runpy.run_path(str(Path(__file__).with_name('test_xml_integrado.py')))
KEY = h['CHAVE']
KEY2 = KEY[:25] + '000000124' + KEY[34:]


def setup(tmp_path):
    cfg = h['helpers']['settings_web'](tmp_path)
    h['helpers']['criar_banco'](cfg.database_path)
    payload = h['xml']().decode()
    first = re.search(r'<det .*?</det>', payload).group()
    second = first.replace('nItem="1"', 'nItem="2"').replace('Produto de teste', 'Aço especial').replace('00007', 'ABC009').replace('01234567', '87654321')
    payload = payload.replace('</det>', '</det>' + second, 1)
    paths = []
    for i, data in enumerate((payload.encode(), h['xml'](chave=KEY2, dest='Outra empresa'))):
        path = tmp_path / f'item{i}.xml'
        path.write_bytes(data)
        paths.append((path, path.name))
    xml_store.import_batch(cfg.database_path, CNPJ_PADRAO, paths)
    return cfg


def test_item_filter_is_literal_and_applies_before_pagination_and_to_all_kinds(tmp_path):
    cfg = setup(tmp_path)
    for query in ('AÇO', 'ABC009', '87654321'):
        for kind in ('notas', 'itens', 'retencoes'):
            report = xml_store.query_report(cfg.database_path, CNPJ_PADRAO, kind=kind, item_query=query)
            assert report['total'] == 1
            assert report['rows'][0]['Chave Acesso'] == KEY
    for query in ('Outra empresa', '%', "' OR 1=1 --"):
        assert xml_store.query_report(cfg.database_path, CNPJ_PADRAO, kind='itens', item_query=query)['total'] == 0


def test_item_screen_and_return_keep_original_filters(tmp_path):
    cfg = setup(tmp_path)
    with TestClient(create_app(cfg)) as client:
        original = '/xml?pagina=2&q=123&situacao=autorizada&data_inicial=01%2F10%2F2026&item=teste'
        page = client.get(original).text
        link = unescape(re.search(r'data-items-open href="([^"]+)"', page).group(1))
        assert '/xml/itens/' + KEY in link
        detail = client.get(link)
        assert detail.status_code == 200
        assert '<h1>ITENS | NF 123</h1>' in detail.text
        assert 'Aço especial' in detail.text and 'Produto de teste' in detail.text
        assert 'Outra empresa' not in detail.text
        back = unescape(re.search(r'href="([^"]+)" data-items-back', detail.text).group(1))
        assert back == original
        filtered = client.get('/xml/itens/' + KEY, params={'item': 'AÇO', 'voltar': original})
        assert 'Aço especial' in filtered.text and 'Produto de teste' not in filtered.text
        assert 'name="chave" value="' + KEY in filtered.text
        assert 'name="item" value="AÇO"' in filtered.text
        assert 'name="voltar" value="' in filtered.text


def test_detail_export_has_only_matching_items_of_selected_note(tmp_path):
    cfg = setup(tmp_path)
    with TestClient(create_app(cfg)) as client:
        params = {'chave': KEY, 'item': 'teste', 'tipo': 'itens'}
        response = client.get('/xml/exportar', params={**params, 'formato': 'csv'})
        rows = list(csv.DictReader(StringIO(response.content.decode('utf-8-sig')), delimiter=';'))
        assert len(rows) == 1 and rows[0]['Chave Acesso'] == KEY
        response = client.get('/xml/exportar', params={**params, 'formato': 'xlsx'})
        wb = load_workbook(BytesIO(response.content))
        assert wb.sheetnames == ['Itens']
        assert wb['Itens'].max_row == 2
        assert client.get('/xml/itens/invalid').status_code == 400
        assert client.get('/xml/itens/' + '0' * 44).status_code == 404
        assert client.get('/xml/itens/' + KEY, params={'item': 'a' * 201}).status_code == 400
        for bad in ('https://evil.example', '//evil.example', '/admin', '/xml\\evil'):
            page = client.get('/xml/itens/' + KEY, params={'voltar': bad}).text
            assert 'href="/xml" data-items-back' in page


def test_item_filter_works_with_encrypted_database(tmp_path):
    from nfe_consulta.seguranca_banco import migrar_banco
    cfg = setup(tmp_path)
    encrypted = tmp_path / 'seguro.db'
    password = 'senha-teste-123456'
    migrar_banco(cfg.database_path, encrypted, password)
    report = xml_store.query_report(encrypted, CNPJ_PADRAO, password=password,
                                    kind='itens', note_key=KEY, item_query='AÇO')
    assert report['total'] == 1 and report['rows'][0]['Descricao'] == 'Aço especial'
