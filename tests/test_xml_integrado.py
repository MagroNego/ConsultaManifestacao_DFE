from datetime import date
from io import BytesIO
from pathlib import Path
import asyncio
import runpy
from zipfile import ZipFile

import pytest
import httpx
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from nfe_consulta.banco import BancoManifestacoes
from nfe_consulta.config import CNPJ_PADRAO
from nfe_consulta.web.app import create_app
from nfe_consulta.web.auth import PUBLIC_USER, WebUser, csrf_token
from nfe_consulta.web import xml_store as store
from nfe_consulta.modelos import Manifestacao, RetornoDistribuicao

helpers = runpy.run_path(str(Path(__file__).with_name('test_web.py')))
CHAVE = '332610' + CNPJ_PADRAO + '55001' + '000000123' + '1' + '00000001' + '0'


def xml(chave=CHAVE, dest='=1+1', emission='2026-10-01', namespace=True):
    ns = 'xmlns="http://www.portalfiscal.inf.br/nfe"' if namespace else ''
    return f'''<nfeProc {ns}><NFe><infNFe Id="NFe{chave}"><ide><mod>55</mod><nNF>{int(chave[25:34])}</nNF><serie>{int(chave[22:25])}</serie><dhEmi>{emission}T10:00:00-03:00</dhEmi></ide><emit><CNPJ>{chave[6:20]}</CNPJ><xNome>Empresa teste</xNome></emit><dest><xNome>{dest}</xNome></dest><det nItem="1"><prod><cProd>00007</cProd><xProd>Produto de teste</xProd><NCM>01234567</NCM><CFOP>5102</CFOP><vUnCom>0.1234567890</vUnCom><qCom>3.0000</qCom><vProd>0.37</vProd></prod><imposto><ICMS><ICMSSN102><orig>0</orig><CSOSN>102</CSOSN></ICMSSN102></ICMS><PIS><PISQtde><CST>03</CST><qBCProd>3</qBCProd><vAliqProd>0.1234</vAliqProd><vPIS>0.37</vPIS></PISQtde></PIS></imposto></det><total><ICMSTot><vNF>0.37</vNF></ICMSTot><retTrib><vRetPIS>1.50</vRetPIS><vRetCOFINS>0.00</vRetCOFINS></retTrib></total></infNFe></NFe></nfeProc>'''.encode()


def setup(tmp_path):
    path = tmp_path / 'historico.db'
    BancoManifestacoes(str(path)).fechar()
    source = tmp_path / 'nota.xml'
    source.write_bytes(xml())
    return path, source


def zip_bytes(entries):
    out = BytesIO()
    with ZipFile(out, 'w') as archive:
        for name, content in entries:
            archive.writestr(name, content)
    return out.getvalue()


def login(client, cfg):
    return csrf_token(PUBLIC_USER, cfg)


def test_arquivo_cumulativo_duplicatas_download_e_reabertura(tmp_path):
    path, source = setup(tmp_path)
    first = store.import_batch(path, CNPJ_PADRAO, [(source, 'mensal.xml')])
    assert first['importadas'] == 1
    assert first['itens'] == 1 and first['retencoes'] == 1
    second = store.import_batch(path, CNPJ_PADRAO, [(source, 'outra.xml')])
    assert second['duplicadas'] == 1 and second['importadas'] == 0
    assert store.download_xml(path, CNPJ_PADRAO, CHAVE) == xml()
    assert store.existing_keys(path, CNPJ_PADRAO, [CHAVE]) == {CHAVE}
    report = store.query_report(path, CNPJ_PADRAO, kind='itens', query='produto')
    assert report['total'] == 1
    row = report['rows'][0]
    assert row['Codigo'] == '00007' and row['CSOSN'] == '102'
    assert row['Vl. Unitario'] == '0,123456789'
    assert row['Base PIS'] == '' and row['Qtd Base PIS'] == '3'
    assert store.query_report(path, CNPJ_PADRAO, inicio=date(2026, 11, 1))['total'] == 0
    history, batch = store.import_history(path, batch_id=first['id'])
    assert len(history) == 2 and batch['importadas'] == 1


def test_zip_invalido_nao_mantem_lote_parcial(tmp_path):
    path, source = setup(tmp_path)
    broken = tmp_path / 'ruim.zip'
    broken.write_bytes(b'nao e zip')
    with pytest.raises(ValueError, match='ZIP inválido'):
        store.import_batch(path, CNPJ_PADRAO, [(source, 'a.xml'), (broken, 'b.zip')])
    assert store.query_report(path, CNPJ_PADRAO)['total'] == 0


def test_zip_xml_invalido_traversal_e_duplicado(tmp_path):
    path, source = setup(tmp_path)
    archive = tmp_path / 'lote.zip'
    archive.write_bytes(zip_bytes([('pasta/a.xml', xml()), ('pasta/b.xml', xml()), ('../fora.xml', xml()), ('evento.xml', b'<resEvento/>'), ('readme.txt', b'teste')]))
    result = store.import_batch(path, CNPJ_PADRAO, [(archive, 'lote.zip')])
    assert (result['importadas'], result['duplicadas'], result['erros']) == (1, 1, 3)
    assert not (tmp_path.parent / 'fora.xml').exists()


@pytest.mark.parametrize('payload', [b'<!DOCTYPE NFe><NFe/>', '<!DOCTYPE NFe><NFe/>'.encode('utf-16'), b'<resNFe/>', b'<procEventoNFe/>', b'<nfeProc>', xml(chave='33261012345678000190550010000001231000000010')])
def test_xml_hostil_incompleto_ou_de_outro_emitente(payload):
    with pytest.raises(Exception):
        store.read_document(payload, CNPJ_PADRAO)


def test_xml_sem_namespace_e_nfe_sem_envelope():
    assert store.read_document(xml(namespace=False), CNPJ_PADRAO)[0]['chave'] == CHAVE
    from defusedxml import ElementTree as ET
    root = ET.fromstring(xml())
    assert store.read_document(ET.tostring(list(root)[0]), CNPJ_PADRAO)[0]['chave'] == CHAVE


def test_web_importa_4000_xmls_soltos(tmp_path):
    cfg = helpers['settings_web'](tmp_path)
    helpers['criar_banco'](cfg.database_path)
    app = create_app(cfg)
    files = [('files', (f'{index}.xml', xml(chave=CHAVE[:25] + f'{index:09d}' + CHAVE[34:]),
                        'application/xml')) for index in range(1, 4001)]
    with TestClient(app) as client:
        token = login(client, cfg)
        response = client.post('/xml/importar', data={'csrf': token}, files=files,
                               follow_redirects=False)
        assert response.status_code == 303
    assert store.query_report(cfg.database_path, CNPJ_PADRAO)['total'] == 4000
    _, batch = store.import_history(cfg.database_path, batch_id=1)
    assert batch['importadas'] == 4000 and batch['erros'] == 0


def test_web_rejeita_mais_de_5000_arquivos_sem_importar(tmp_path):
    cfg = helpers['settings_web'](tmp_path)
    helpers['criar_banco'](cfg.database_path)
    app = create_app(cfg)
    with TestClient(app) as client:
        token = login(client, cfg)
        files = [('files', (f'{index}.xml', b'', 'application/xml')) for index in range(5001)]
        response = client.post('/xml/importar', data={'csrf': token}, files=files)
        assert response.status_code == 400
        assert '5000' in response.json()['detail']
    assert store.query_report(cfg.database_path, CNPJ_PADRAO)['total'] == 0


def test_excel_csv_precisao_formula_e_codigos(tmp_path):
    path, source = setup(tmp_path)
    store.import_batch(path, CNPJ_PADRAO, [(source, 'a.xml')])
    report = store.query_report(path, CNPJ_PADRAO, kind='itens', export=True)
    assert "'=1+1" in store.export_csv(report).decode('utf-8-sig')
    notes = store.query_report(path, CNPJ_PADRAO, kind='notas', export=True)
    wb = load_workbook(BytesIO(store.export_excel([('Notas', notes), ('Itens', report)])))
    sheet = wb['Itens']
    heads = [c.value for c in sheet[1]]
    assert sheet.cell(2, heads.index('Destinatario Razao') + 1).data_type == 's'
    assert sheet.cell(2, heads.index('Chave Acesso') + 1).value == CHAVE
    assert sheet.cell(2, heads.index('Codigo') + 1).value == '00007'
    unit = sheet.cell(2, heads.index('Vl. Unitario') + 1)
    assert unit.value == pytest.approx(.1235)
    assert unit.number_format == '#,##0.0000'
    quantity = sheet.cell(2, heads.index('Qtd') + 1)
    assert quantity.value == 3 and quantity.number_format == '#,##0'
    assert store.EXCEL_EXCLUDED_COLUMNS.isdisjoint(heads)
    assert all(key in report['columns'] for key in store.EXCEL_EXCLUDED_COLUMNS)
    for title in ('Notas', 'Itens'):
        current = wb[title]
        columns = [cell.value for cell in current[1]]
        emission = current.cell(2, columns.index('Data Emissao') + 1)
        assert emission.value.date() == date(2026, 10, 1)
        assert emission.number_format == 'dd/mm/yyyy'


def test_web_lote_manifestacao_download_export_e_permissoes(tmp_path, monkeypatch):
    cfg = helpers['settings_web'](tmp_path)
    helpers['criar_banco'](cfg.database_path)
    def proibido(*args, **kwargs):
        pytest.fail('Importação/consulta XML tentou consultar SEFAZ')
    monkeypatch.setattr('nfe_consulta.web.app.sincronizar_configurado', proibido)
    app = create_app(cfg)
    with TestClient(app) as client:
        assert client.get('/xml',follow_redirects=False).status_code == 200
        token = login(client, cfg)
        assert client.post('/xml/importar', data={'csrf':'errado'}, files={'files':('a.xml',xml())}).status_code == 403
        response = client.post('/xml/importar', data={'csrf':token}, files={'files':('mensal.zip',zip_bytes([('a.xml',xml())]),'application/zip')})
        assert response.status_code == 200 and '1 importada(s)' in response.text
        assert '=1+1' in client.get('/xml').text
        assert client.get('/xml/arquivo/'+CHAVE).content == xml()
        assert client.get('/xml/arquivo/'+'9'*44).status_code == 404
        assert client.get('/xml/exportar?formato=csv&tipo=itens&q=produto').status_code == 200
        data = client.get('/xml/exportar?formato=xlsx').content
        assert load_workbook(BytesIO(data)).sheetnames == ['Notas','Itens','Retencoes']
        assert client.get('/xml/lote/1/registro').status_code == 200
        banco = BancoManifestacoes(str(cfg.database_path))
        banco.salvar_retorno(CNPJ_PADRAO, RetornoDistribuicao(138,'ok','1'.zfill(15),'1'.zfill(15),((CHAVE,Manifestacao('210240','Operação não Realizada','2026-10-01T10:00:00-03:00','1')),),0))
        banco.fechar()
        consulta = client.get('/consulta',params={'chave':CHAVE})
        assert 'Baixar XML' in consulta.text and '/xml/arquivo/'+CHAVE in consulta.text
        assert client.get('/xml?tipo=itens&q=produto').status_code == 200
    with TestClient(create_app(cfg)) as client:
        assert client.get('/xml/arquivo/'+CHAVE).status_code == 200
        login(client,cfg)
        assert '1 nota arquivada' in client.get('/xml').text
        assert client.get('/xml/arquivo/'+CHAVE).content == xml()
        assert client.get('/xml/lote/1/registro').status_code == 200


def test_limites_zip_e_xml(tmp_path, monkeypatch):
    path, source = setup(tmp_path)
    monkeypatch.setattr(store, 'MAX_XML', 20)
    report = store.import_batch(path, CNPJ_PADRAO, [(source,'a.xml')])
    assert report['erros'] == 1 and report['importadas'] == 0
    monkeypatch.setattr(store,'MAX_EXPANDED', 5)
    with pytest.raises(ValueError,match='descompactados'):
        store.import_batch(path, CNPJ_PADRAO, [(source,'a.xml')])


def test_xml_criptografado_nao_vaza_banco_plaintext(tmp_path):
    from nfe_consulta.seguranca_banco import migrar_banco
    path, source = setup(tmp_path)
    secure = tmp_path / 'seguro.db'
    password = 'SenhaTeste-SQLCipher-123!'
    migrar_banco(path, secure, password)
    store.import_batch(secure,CNPJ_PADRAO,[(source,'a.xml')],password=password)
    assert store.download_xml(secure,CNPJ_PADRAO,CHAVE,password=password) == xml()
    assert b'Produto de teste' not in secure.read_bytes()
    assert store.query_report(path,CNPJ_PADRAO)['total'] == 0


def test_numero_com_expoente_extremo_e_recusado():
    with pytest.raises(ValueError, match='numérico'):
        store.read_document(xml().replace(b'3.0000', b'1E+10000000'), CNPJ_PADRAO)


def test_upload_sem_content_length_respeita_limite_de_corpo(tmp_path, monkeypatch):
    cfg = helpers['settings_web'](tmp_path)
    helpers['criar_banco'](cfg.database_path)
    monkeypatch.setattr('nfe_consulta.web.xml_upload_limit.MAX_UPLOAD', 100)
    app = create_app(cfg)
    with TestClient(app) as client:
        token = login(client, cfg)
        cookies = dict(client.cookies)
    async def enviar():
        async def chunks():
            yield (f'--qa\r\nContent-Disposition: form-data; name="csrf"\r\n\r\n{token}\r\n--qa\r\nContent-Disposition: form-data; name="files"; filename="a.xml"\r\nContent-Type: application/xml\r\n\r\n').encode()
            for _ in range(34):
                yield b'x' * (64*1024)
            yield b'\r\n--qa--\r\n'
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://testserver', cookies=cookies) as client:
            return await client.post('/xml/importar', content=chunks(), headers={'Content-Type':'multipart/form-data; boundary=qa'})
    response = asyncio.run(enviar())
    assert response.status_code == 413
    assert store.query_report(cfg.database_path,CNPJ_PADRAO)['total'] == 0


def test_meses_cumulativos_paginacao_e_exportacao_integral(tmp_path):
    path, source = setup(tmp_path)
    store.import_batch(path, CNPJ_PADRAO, [(source, 'outubro.xml')])
    september = tmp_path / 'setembro.zip'
    keys = [CHAVE[:4] + '09' + CHAVE[6:25] + str(i).zfill(9) + CHAVE[34:] for i in range(1, 102)]
    september.write_bytes(zip_bytes([(f'{i}.xml', xml(chave=key, emission='2026-09-15')) for i, key in enumerate(keys)]))
    assert store.import_batch(path, CNPJ_PADRAO, [(september,'setembro.zip')])['importadas'] == 101
    report = store.query_report(path, CNPJ_PADRAO)
    assert report['notas'] == 102 and report['pages'] == 2 and len(report['rows']) == 100
    assert len(store.query_report(path, CNPJ_PADRAO, page=2)['rows']) == 2
    filtered = store.query_report(path, CNPJ_PADRAO, inicio=date(2026,9,1), fim=date(2026,9,30), export=True)
    assert filtered['total'] == len(filtered['rows']) == 101
    source.write_bytes(xml(dest='Conteúdo alterado'))
    assert store.import_batch(path, CNPJ_PADRAO, [(source,'repetida.xml')])['duplicadas'] == 1
    assert store.download_xml(path,CNPJ_PADRAO,CHAVE) == xml()


def test_importacao_travada_e_texto_html_escapado(tmp_path):
    cfg = helpers['settings_web'](tmp_path)
    helpers['criar_banco'](cfg.database_path)
    app = create_app(cfg)
    with TestClient(app) as client:
        token = login(client, cfg)
        app.state.sync_lock.acquire()
        try:
            response = client.post('/xml/importar', data={'csrf':token}, files={'files':('a.xml',xml())})
            assert response.status_code == 409
            assert store.query_report(cfg.database_path,CNPJ_PADRAO)['total'] == 0
        finally:
            app.state.sync_lock.release()
        response = client.post('/xml/importar', data={'csrf':token}, files={'files':('a.xml',xml(dest='&lt;script&gt;alert(1)&lt;/script&gt;'))})
        assert response.status_code == 200
        assert '<script>alert(1)</script>' not in response.text
        assert '&lt;script&gt;alert(1)&lt;/script&gt;' in client.get('/xml').text


@pytest.mark.parametrize('number', [b'0E-10000000', b'1E+10000000'])
def test_total_nota_com_expoente_extremo_recusado(number):
    with pytest.raises(ValueError, match='numérico'):
        store.read_document(xml().replace(b'<vNF>0.37</vNF>', b'<vNF>'+number+b'</vNF>'), CNPJ_PADRAO)


@pytest.mark.parametrize('emission, expected', [('2026-05-10', '10/05/2026'), ('2026-10-05', '05/10/2026')])
def test_datas_xml_em_formato_brasileiro_sem_inverter_dia_mes(tmp_path, emission, expected):
    cfg = helpers['settings_web'](tmp_path)
    helpers['criar_banco'](cfg.database_path)
    source = tmp_path / 'nota.xml'
    source.write_bytes(xml(emission=emission))
    store.import_batch(cfg.database_path, CNPJ_PADRAO, [(source, 'nota.xml')])
    with TestClient(create_app(cfg)) as client:
        for tipo in ['notas', 'itens', 'retencoes']:
            response = client.get('/xml', params={'tipo': tipo})
            assert response.status_code == 200
            assert expected in response.text
            assert emission not in response.text
    assert store.query_report(cfg.database_path, CNPJ_PADRAO)['rows'][0]['Data Emissao'] == emission
