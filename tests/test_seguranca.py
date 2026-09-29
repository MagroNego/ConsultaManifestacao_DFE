import base64
import csv
import gzip
import io

import pytest
from openpyxl import load_workbook

from nfe_consulta.config import NOME_PLANILHA
from nfe_consulta.csv_writer import gravar_csv
from nfe_consulta.modelos import NfeErroResposta, ResultadoConsulta
from nfe_consulta.parser_distribuicao import MAX_XML_BYTES, parse_retorno_distribuicao
from nfe_consulta.xlsx_writer import gravar_xlsx


def test_rejeita_entidades_xml():
    xml = ('<!DOCTYPE retDistDFeInt [<!ENTITY seg "expandir">]>'
           '<retDistDFeInt><cStat>137</cStat><xMotivo>&seg;</xMotivo>'
           '<ultNSU>0</ultNSU><maxNSU>0</maxNSU></retDistDFeInt>')
    with pytest.raises(NfeErroResposta):
        parse_retorno_distribuicao(xml)


def test_bloqueia_gzip_que_expande_demais():
    zipado = base64.b64encode(gzip.compress(b" " * (MAX_XML_BYTES + 1))).decode()
    xml = ('<retDistDFeInt><cStat>138</cStat><ultNSU>1</ultNSU><maxNSU>1</maxNSU>'
           f'<docZip NSU="1" schema="evento">{zipado}</docZip></retDistDFeInt>')
    with pytest.raises(NfeErroResposta, match="limite"):
        parse_retorno_distribuicao(xml)


def test_limite_de_doczip_inclui_resumos():
    resumo = gzip.compress(b"<resNFe><chNFe>" + b"1" * 44 + b"</chNFe></resNFe>")
    zipado = base64.b64encode(resumo).decode()
    xml = ('<retDistDFeInt><cStat>138</cStat><ultNSU>1</ultNSU><maxNSU>1</maxNSU>'
           + ''.join(f'<docZip NSU="{i}" schema="resNFe_v1.01.xsd">{zipado}</docZip>' for i in range(51))
           + '</retDistDFeInt>')
    with pytest.raises(NfeErroResposta, match="50 documentos"):
        parse_retorno_distribuicao(xml)


def test_exportacoes_tratam_texto_que_parece_formula(tmp_path):
    resultado = ResultadoConsulta(
        '=HYPERLINK("https://example.com")',
        0,
        "",
        None,
        (),
        "=SUM(1,1)",
    )
    destino = tmp_path / NOME_PLANILHA
    gravar_xlsx(str(destino), [resultado], "=SUM(2,2)")
    ws = load_workbook(destino).active
    assert ws.title == "Consulta_Manifestacao_YAB"
    assert ws["A5"].data_type == "s"
    assert ws["J5"].data_type == "s"
    assert ws["A3"].value is None

    texto = io.StringIO()
    gravar_csv(texto, [resultado], "=SUM(2,2)")
    registro = next(csv.DictReader(io.StringIO(texto.getvalue()), delimiter=";"))
    assert registro["chave"].startswith("'=")
    assert registro["erro"].startswith("'=")
    assert registro["cobertura"].startswith("'=")
