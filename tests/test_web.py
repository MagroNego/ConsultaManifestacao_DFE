from conftest import authenticated_user
from io import BytesIO
from types import SimpleNamespace
import asyncio
import time
import json
import pytest

from fastapi.testclient import TestClient
import httpx
from openpyxl import load_workbook

from nfe_consulta.banco import BancoManifestacoes
from nfe_consulta.modelos import InformacaoNota, Manifestacao, RetornoDistribuicao
from nfe_consulta.web.app import create_app
from nfe_consulta.web.auth import PUBLIC_USER, WebUser, csrf_token
from nfe_consulta.web.settings import WebSettings


CHAVE = "33260812345678000199550010000917791147439711"


def settings_web(tmp_path, *, admin=False, auth_mode="dev", admin_users=frozenset()):
    return WebSettings(
        local_auth_enabled=True,
        environment="development",
        auth_mode=auth_mode,
        user_header="X-NFE-User",
        name_header="X-NFE-Name",
        proxy_secret_header="X-NFE-Proxy-Secret",
        proxy_secret="segredo-de-proxy-com-tamanho-suficiente",
        admin_users=admin_users,
        dev_user="usuario@empresa.local",
        dev_name="Usuário Teste",
        dev_admin=admin,
        csrf_secret="c" * 64,
        database_path=tmp_path / "dados" / "historico.db",
        database_password=None,
        certificate_store="CurrentUser",
        certificate_thumbprint=None,
        bind_host="127.0.0.1",
        bind_port=8080,
        root_path="",
        forwarded_allow_ips="127.0.0.1",
        audit_log=tmp_path / "logs" / "web_audit.log",
        sync_cooldown_minutes=60,
        admin_username="admin",
        admin_password="Senha-Admin-123!",
        admin_session_minutes=30,
        admin_cookie_name="nfe_admin_session",
        certificate_path_file=tmp_path / "secrets" / "cert-path.txt",
        certificate_password_file=tmp_path / "secrets" / "cert-password.txt",
        database_path_file=tmp_path / "secrets" / "db-path.txt",
        email_recipients_file=tmp_path / "secrets" / "email-recipients.json",
        smtp_password_file=tmp_path / "secrets" / "smtp-password.txt",
    )


def criar_banco(caminho):
    caminho.parent.mkdir(parents=True, exist_ok=True)
    banco = BancoManifestacoes(str(caminho))
    banco.fechar()


def entrar_admin(client, cfg):
    client.cookies.clear()
    response = client.post('/admin/login', data={'csrf': csrf_token(PUBLIC_USER, cfg),
        'username': cfg.admin_username, 'password': cfg.admin_password}, follow_redirects=False)
    assert response.status_code == 303
    user = WebUser(cfg.admin_username, cfg.admin_username, True, client.cookies[cfg.admin_cookie_name], 'admin')
    return csrf_token(user, cfg)


def test_admin_cadastra_alertas_e_bloqueia_csrf_invalido(tmp_path):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    app = create_app(cfg)
    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        publico = client.post(
            "/atualizar/alertas-email",
            data={"csrf": token, "recipients": "a@empresa.com.br"},
            follow_redirects=False,
        )
        assert publico.status_code == 303
        assert cfg.email_recipients_file.exists()

        pagina = client.get("/atualizar")
        assert "Alertas por e-mail" in pagina.text
        assert "Sem criptografia" in pagina.text
        invalido = client.post("/atualizar/alertas-email", data={
            "csrf": "invalido", "recipients": "a@empresa.com.br",
        })
        assert invalido.status_code == 403
        salvo = client.post("/atualizar/alertas-email", data={
            "csrf": token,
            "recipients": "a@empresa.com.br\nb@empresa.com.br",
        }, follow_redirects=False)
        assert salvo.status_code == 303
        assert "a@empresa.com.br" in client.get("/atualizar").text


def test_excel_web_gera_xlsx_com_banco_local(tmp_path):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    app = create_app(cfg)
    token = csrf_token(PUBLIC_USER, cfg)

    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.post(
            "/excel",
            data={"csrf": token},
            files={
                "files": (
                    "CHAVES.txt",
                    CHAVE + "\n",
                    "text/plain",
                )
            },
        )

    assert resposta.status_code == 200
    assert resposta.content.startswith(b"PK")
    assert "spreadsheetml.sheet" in resposta.headers["content-type"]
    assert cfg.audit_log.is_file()


def test_excel_web_bloqueia_csrf_invalido(tmp_path):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    app = create_app(cfg)

    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.post(
            "/excel",
            data={"csrf": "invalido"},
            files={"files": ("CHAVES.txt", CHAVE + "\n", "text/plain")},
        )

    assert resposta.status_code == 403


def test_admin_pode_disparar_atualizacao_sefaz(tmp_path, monkeypatch):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    app = create_app(cfg)

    chamadas = []

    def fake_sync(parametros, **kwargs):
        chamadas.append(parametros)
        return SimpleNamespace(
            lotes=1,
            eventos_novos=2,
            ult_nsu="10".zfill(15),
            max_nsu="10".zfill(15),
            completo=True,
            cache=False,
        )

    monkeypatch.setattr("nfe_consulta.web.sync_runtime.sincronizar_banco", fake_sync)

    with TestClient(app) as client:
        token = entrar_admin(client, cfg)

        resposta = client.post(
            "/atualizar/sincronizar",
            data={"csrf": token, "max_lotes": "25", "confirmacao": "sincronizar"},
            follow_redirects=False,
        )
        history = client.get("/status")

    assert resposta.status_code == 303
    assert "Últimas sincronizações" in history.text
    assert "Manual" in history.text
    assert "Execução concluída" in history.text
    assert ">2</td>" in history.text
    assert chamadas[0].max_lotes == 25
    assert chamadas[0].cooldown_minutos == 60


@pytest.mark.parametrize("pausa", [False, True])
def test_post_direto_durante_bloqueio_nao_consulta_sefaz(tmp_path, monkeypatch, pausa):
    from nfe_consulta.config import CNPJ_PADRAO
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    banco = BancoManifestacoes(str(cfg.database_path))
    if pausa:
        banco.pausar_distribuicao(CNPJ_PADRAO, "656")
    else:
        banco.reservar_sincronizacao(CNPJ_PADRAO, 60)
    banco.fechar()
    def proibido(*args, **kwargs):
        pytest.fail("POST durante bloqueio iniciou sincronização")
    monkeypatch.setattr("nfe_consulta.web.app.sincronizar_configurado", proibido)
    banco_anterior = cfg.database_path.read_bytes()
    with TestClient(create_app(cfg)) as client:
        token = entrar_admin(client, cfg)
        assert "data-sync-confirm" not in client.get("/atualizar").text
        resposta = client.post("/atualizar/sincronizar", data={
            "csrf": token, "max_lotes": "50", "confirmacao": "sincronizar",
            "can_sync": "true", "enabled": "true", "force": "1",
        })
    assert resposta.status_code == 429
    assert "Sincronização bloqueada temporariamente" in resposta.text

    assert cfg.database_path.read_bytes() == banco_anterior


@pytest.mark.parametrize("pausa", [False, True])
def test_agendamento_durante_bloqueio_nao_consulta_sefaz(tmp_path, monkeypatch, pausa):
    from nfe_consulta.config import CNPJ_PADRAO
    from nfe_consulta.web.scheduler import executar_sincronizacao_automatica
    from dataclasses import replace
    cfg = replace(settings_web(tmp_path), auto_sync_enabled=True)
    criar_banco(cfg.database_path)
    banco = BancoManifestacoes(str(cfg.database_path))
    if pausa:
        banco.pausar_distribuicao(CNPJ_PADRAO, "656")
    else:
        banco.reservar_sincronizacao(CNPJ_PADRAO, 60)
    banco.fechar()
    def proibido(*args, **kwargs):
        pytest.fail("Agendamento bloqueado chegou ao certificado ou à SEFAZ")
    monkeypatch.setattr("nfe_consulta.servico.resolver_certificado", proibido)
    monkeypatch.setattr("nfe_consulta.servico.sincronizar", proibido)
    app = create_app(cfg)
    asyncio.run(executar_sincronizacao_automatica(app))
    registro = json.loads(cfg.audit_log.read_text().split(" ", 2)[2])
    assert registro["reason"] == "cooldown"
    assert registro["result"] == "ignorado"
    assert app.state.sync_lock.acquire(blocking=False)
    app.state.sync_lock.release()


def test_sincronizacao_manual_exige_confirmacao_no_servidor(tmp_path, monkeypatch):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    def proibido(*args, **kwargs):
        pytest.fail("Sincronização sem confirmação chegou à SEFAZ")
    monkeypatch.setattr("nfe_consulta.web.app.sincronizar_configurado", proibido)
    with TestClient(create_app(cfg)) as client:
        token = entrar_admin(client, cfg)
        resposta = client.post("/atualizar/sincronizar", data={"csrf": token})
    assert resposta.status_code == 400
    assert "Confirme" in resposta.text


def test_status_mostra_ultima_gravacao_sem_chamar_sefaz(tmp_path):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    app = create_app(cfg)

    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.get("/status")

    assert resposta.status_code == 200
    assert "Última gravação" in resposta.text
    assert "Sem sincronização" in resposta.text


def test_headers_de_seguranca_e_healthcheck(tmp_path):
    cfg = settings_web(tmp_path)
    app = create_app(cfg)

    with TestClient(app) as client:
        token_usuario = authenticated_user(client)
        health = client.get("/healthz")
        home = client.get("/")

    assert health.status_code == 200
    assert health.json()["status"] == "ok"
    assert home.headers["x-frame-options"] == "DENY"
    assert home.headers["x-content-type-options"] == "nosniff"
    assert "default-src 'self'" in home.headers["content-security-policy"]


def test_tela_consulta_tem_filtros_e_fluxo_por_txt(tmp_path):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    app = create_app(cfg)

    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.get("/")

    assert resposta.status_code == 200
    assert "CHAVES.txt" in resposta.text
    assert "webkitdirectory" not in resposta.text
    assert ">Pasta<" not in resposta.text
    assert "Data inicial do evento" in resposta.text
    assert "Data final do evento" in resposta.text
    assert "Manifestação" in resposta.text
    assert "Consulta por arquivo" in resposta.text


def test_consulta_rapida_por_numero_mostra_eventos_do_banco(tmp_path, monkeypatch):
    cfg = settings_web(tmp_path)
    banco = BancoManifestacoes(str(cfg.database_path))
    evento = Manifestacao(
        codigo="210210",
        descricao="Ciência da Operação",
        data="2026-09-25T08:00:00-03:00",
        protocolo="135260000000001",
        nsu="563664",
        schema="procEventoNFe_v1.00.xsd",
    )
    banco.salvar_retorno(
        "16840128000101",
        RetornoDistribuicao(
            status_codigo=138,
            status_motivo="Documento localizado",
            ult_nsu="563664".zfill(15),
            max_nsu="563664".zfill(15),
            manifestacoes=((CHAVE, evento),),
        ),
    )
    banco.fechar()

    monkeypatch.setattr(
        "nfe_consulta.web.sync_runtime.sincronizar_banco",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("Consulta rápida não deve acessar a SEFAZ")
        ),
    )

    app = create_app(cfg)
    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.get("/consulta", params={"numero": "91779"})

    assert resposta.status_code == 200
    assert "91779" in resposta.text
    assert "Ciência da Operação" in resposta.text
    assert "135260000000001" in resposta.text
    assert ">1<" in resposta.text


def test_consulta_rapida_nao_aceita_expressao_sql(tmp_path):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    app = create_app(cfg)

    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.get("/consulta", params={"numero": "91779 OR 1=1"})

    assert resposta.status_code == 200
    assert "Informe somente números no campo Número da NF." in resposta.text


def test_cabecalho_usa_logos_yorozu_por_tema(tmp_path):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    app = create_app(cfg)

    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.get("/")
        logo_clara = client.get("/static/img/yorozu-light.png")
        logo_escura = client.get("/static/img/yorozu-dark.png")

    assert resposta.status_code == 200
    assert "/static/img/yorozu-light.png" in resposta.text
    assert "/static/img/yorozu-dark.png" in resposta.text
    assert logo_clara.status_code == 200
    assert logo_escura.status_code == 200
    assert logo_clara.headers["content-type"] == "image/png"
    assert logo_escura.headers["content-type"] == "image/png"


def test_admin_configura_certificado_por_caminho_no_servidor(tmp_path, monkeypatch):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    app = create_app(cfg)

    cert_path = tmp_path / "TI" / "Yorozu.pfx"
    cert_path.parent.mkdir(parents=True)
    cert_path.write_bytes(b"certificado-sintetico")

    cert = SimpleNamespace(
        thumbprint="ABC",
        subject="CN=YOROZU:16840128000101",
        issuer="ICP-Brasil",
        valid_to="2027-07-10",
        cnpj="16840128000101",
    )

    monkeypatch.setattr(
        "nfe_consulta.web.app.carregar_certificado_arquivo",
        lambda *_args, **_kwargs: cert,
    )
    monkeypatch.setattr(
        "nfe_consulta.certificado_config.carregar_certificado_arquivo",
        lambda *_args, **_kwargs: cert,
    )

    with TestClient(app) as client:
        token = entrar_admin(client, cfg)

        resposta = client.post(
            "/atualizar/certificado",
            data={
                "csrf": token,
                "certificate_path": str(cert_path),
                "certificate_password": "Senha-PFX-123!",
            },
            follow_redirects=False,
        )

    assert resposta.status_code == 303
    assert "cert_ok=1" in resposta.headers["location"]
    assert cfg.certificate_path_file.read_text(encoding="utf-8") == str(cert_path.resolve())
    assert cfg.certificate_password_file.read_text(encoding="utf-8") == "Senha-PFX-123!"
    assert not (cfg.certificate_path_file.parent / cert_path.name).exists()


def test_sincronizacao_web_usa_certificado_configurado_em_arquivo(tmp_path, monkeypatch):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)

    cert_path = tmp_path / "TI" / "Yorozu.pfx"
    cert_path.parent.mkdir(parents=True)
    cert_path.write_bytes(b"certificado-sintetico")
    cfg.certificate_path_file.parent.mkdir(parents=True, exist_ok=True)
    cfg.certificate_path_file.write_text(str(cert_path), encoding="utf-8")
    cfg.certificate_password_file.write_text("Senha-PFX-123!", encoding="utf-8")

    app = create_app(cfg)
    chamadas = []

    def fake_sync(parametros, **kwargs):
        chamadas.append(parametros)
        return SimpleNamespace(
            lotes=1,
            eventos_novos=0,
            ult_nsu="10".zfill(15),
            max_nsu="10".zfill(15),
            completo=True,
            cache=False,
        )

    monkeypatch.setattr("nfe_consulta.web.sync_runtime.sincronizar_banco", fake_sync)
    monkeypatch.setattr(
        "nfe_consulta.web.app.carregar_certificado_arquivo",
        lambda *_args, **_kwargs: SimpleNamespace(
            subject="CN=YOROZU",
            valid_to="2027-07-10",
        ),
    )

    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.post(
            "/atualizar/sincronizar",
            data={"csrf": token, "max_lotes": "5", "confirmacao": "sincronizar"},
            follow_redirects=False,
        )

    assert resposta.status_code == 303
    assert chamadas[0].cert_arquivo == cert_path.resolve()
    assert chamadas[0].cert_senha_arquivo == "Senha-PFX-123!"
    assert chamadas[0].cert_thumbprint is None


def test_admin_configura_banco_por_caminho_e_aplicacao_passa_a_usar(tmp_path):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)

    banco_novo = tmp_path / "TI" / "dados" / "central.db"
    banco_novo.parent.mkdir(parents=True, exist_ok=True)
    banco = BancoManifestacoes(str(banco_novo))
    evento = Manifestacao(
        codigo="210210",
        descricao="Ciência da Operação",
        data="2026-09-25T08:00:00-03:00",
        protocolo="135260000000099",
        nsu="563700",
        schema="procEventoNFe_v1.00.xsd",
    )
    banco.salvar_retorno(
        "16840128000101",
        RetornoDistribuicao(
            status_codigo=138,
            status_motivo="Documento localizado",
            ult_nsu="563700".zfill(15),
            max_nsu="563700".zfill(15),
            manifestacoes=((CHAVE, evento),),
        ),
    )
    banco.fechar()

    app = create_app(cfg)
    with TestClient(app) as client:
        token = entrar_admin(client, cfg)

        resposta = client.post(
            "/atualizar/banco",
            data={
                "csrf": token,
                "database_path": str(banco_novo),
            },
            follow_redirects=False,
        )
        assert resposta.status_code == 303
        assert "db_ok=1" in resposta.headers["location"]

        consulta = client.get("/consulta", params={"numero": "91779"})

    assert cfg.database_path_file.read_text(encoding="utf-8") == str(banco_novo.resolve())
    assert consulta.status_code == 200
    assert "135260000000099" in consulta.text


def test_admin_rejeita_arquivo_que_nao_e_banco_de_manifestacoes(tmp_path):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    invalido = tmp_path / "qualquer.db"
    invalido.write_bytes(b"nao-e-um-banco")

    app = create_app(cfg)
    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.post(
            "/atualizar/banco",
            data={
                "csrf": token,
                "database_path": str(invalido),
            },
        )

    assert resposta.status_code == 400
    assert not cfg.database_path_file.exists()


def test_consulta_completa_filtra_por_data_e_manifestacao(tmp_path):
    cfg = settings_web(tmp_path)
    banco = BancoManifestacoes(str(cfg.database_path))

    chave_2 = CHAVE[:25] + "000091780" + CHAVE[34:]
    eventos = (
        (
            CHAVE,
            Manifestacao(
                codigo="210210",
                descricao="Ciência da Operação",
                data="2026-09-10T08:00:00-03:00",
                protocolo="135260000000101",
                nsu="563701",
                schema="procEventoNFe_v1.00.xsd",
            ),
        ),
        (
            chave_2,
            Manifestacao(
                codigo="210240",
                descricao="Operação não Realizada",
                data="2026-09-20T09:30:00-03:00",
                protocolo="135260000000102",
                nsu="563702",
                schema="procEventoNFe_v1.00.xsd",
            ),
        ),
    )
    banco.salvar_retorno(
        "16840128000101",
        RetornoDistribuicao(
            status_codigo=138,
            status_motivo="Documentos localizados",
            ult_nsu="563702".zfill(15),
            max_nsu="563702".zfill(15),
            manifestacoes=eventos,
            informacoes_notas=(InformacaoNota(chave_2, "Fornecedor Teste", True),),
        ),
    )
    banco.fechar()

    app = create_app(cfg)
    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.get(
            "/consulta",
            params={
                "consultar": "1",
                "data_inicial": "2026-09-15",
                "data_final": "2026-09-30",
                "codigo": "210240",
            },
        )

    assert resposta.status_code == 200
    assert "Operação não Realizada" in resposta.text
    assert "Fornecedor Teste" not in resposta.text  # Emitente removido da tabela de consulta.
    assert "Cancelada" in resposta.text
    assert 'name="canceladas"' not in resposta.text
    assert 'name="serie"' not in resposta.text
    assert "91780" in resposta.text
    assert "135260000000101" not in resposta.text
    assert ">91779<" not in resposta.text


def test_exportacao_da_consulta_respeita_os_mesmos_filtros(tmp_path):
    cfg = settings_web(tmp_path)
    banco = BancoManifestacoes(str(cfg.database_path))

    chave_2 = CHAVE[:25] + "000091780" + CHAVE[34:]
    banco.salvar_retorno(
        "16840128000101",
        RetornoDistribuicao(
            status_codigo=138,
            status_motivo="Documentos localizados",
            ult_nsu="563704".zfill(15),
            max_nsu="563704".zfill(15),
            manifestacoes=(
                (
                    CHAVE,
                    Manifestacao(
                        codigo="210210",
                        descricao="Ciência da Operação",
                        data="2026-09-10T08:00:00-03:00",
                        protocolo="135260000000201",
                        nsu="563703",
                        schema="procEventoNFe_v1.00.xsd",
                    ),
                ),
                (
                    chave_2,
                    Manifestacao(
                        codigo="210240",
                        descricao="Operação não Realizada",
                        data="2026-09-20T09:30:00-03:00",
                        protocolo="135260000000202",
                        nsu="563704",
                        schema="procEventoNFe_v1.00.xsd",
                    ),
                ),
            ),
            informacoes_notas=(InformacaoNota(chave_2, "Fornecedor Teste", True),),
        ),
    )
    banco.fechar()

    app = create_app(cfg)
    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.get(
            "/consulta/exportar",
            params={
                "data_inicial": "2026-09-15",
                "data_final": "2026-09-30",
                "codigo": "210240",
            },
        )

    assert resposta.status_code == 200
    assert resposta.content.startswith(b"PK")
    wb = load_workbook(BytesIO(resposta.content), read_only=True)
    ws = wb["Manifestacoes"]
    assert ws["A2"].value == 91780
    assert ws["E2"].value == "Operação não Realizada"
    assert ws["J2"].value == "Fornecedor Teste"
    assert ws.max_column == 12
    assert ws['K1'].value == 'Situação'
    assert "Cancelamento" not in [c.value for c in ws[1]]
    assert ws["A3"].value is None


def test_consulta_identifica_banco_sem_sincronizacao(tmp_path):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    app = create_app(cfg)

    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.get("/consulta", params={"numero": "91779"})

    assert resposta.status_code == 200
    assert "Banco não sincronizado" in resposta.text
    assert "ainda não possui uma sincronização concluída" in resposta.text
    assert "informe a senha" not in resposta.text.casefold()
    assert "sqlcipher" not in resposta.text.casefold()


def test_consulta_publica_nao_expoe_erro_de_senha_do_banco(tmp_path, monkeypatch):
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)

    def falha_banco(_settings):
        raise ValueError("Banco criptografado: informe a senha para abri-lo.")

    monkeypatch.setattr("nfe_consulta.web.app._status_web", falha_banco)

    app = create_app(cfg)
    with TestClient(app) as client:
        token = entrar_admin(client, cfg)
        resposta = client.get("/consulta", params={"numero": "91779"})

    assert resposta.status_code == 200
    assert "Banco não sincronizado" in resposta.text
    assert "informe a senha" not in resposta.text.casefold()
    assert "criptografado" not in resposta.text.casefold()


def test_erro_manual_exibe_diagnostico_no_status_sem_mensagem_bruta(tmp_path, monkeypatch):
    from nfe_consulta.modelos import NfeConsumoIndevidoErro
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    exc = NfeConsumoIndevidoErro('senha=segredo mensagem interna', ult_nsu='25', ult_nsu_enviado='20')
    exc.sync_diagnostics = dict(execucao='b'*32, etapa='distNSU', nsu_antes='10', nsu_salvo='20',
        lotes_gravados=1, respostas_pendentes=0, proxima_tentativa='2026-10-14T22:00:00-03:00')
    def falhar(*args, **kwargs):
        raise exc
    monkeypatch.setattr('nfe_consulta.web.app.sincronizar_configurado', falhar)
    with TestClient(create_app(cfg)) as client:
        token = entrar_admin(client, cfg)
        resposta = client.post('/atualizar/sincronizar', data={
            'csrf': token, 'max_lotes': '50', 'confirmacao': 'sincronizar'})
        history = client.get('/status')
    assert resposta.status_code == 429
    assert 'segredo' not in resposta.text and 'segredo' not in history.text
    assert 'b'*32 in resposta.text and 'b'*32 in history.text
    assert 'NSU salvo: 000000000000020' in history.text
    assert 'Lotes gravados: 1' in history.text
    assert '14/10/2026 22:00 (Brasília)' in history.text
    assert '<details class="card status-summary sync-logs">' in history.text


def test_post_fora_do_horario_registra_janela_em_vez_de_cooldown(tmp_path, monkeypatch):
    from datetime import datetime
    from nfe_consulta.janela_consulta import BRASILIA
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    monkeypatch.setattr('nfe_consulta.janela_consulta.agora_brasilia',
                        lambda: datetime(2026, 10, 13, 18, tzinfo=BRASILIA))
    def proibido(*args, **kwargs):
        pytest.fail('Consulta fora da janela')
    monkeypatch.setattr('nfe_consulta.web.app.sincronizar_configurado', proibido)
    with TestClient(create_app(cfg)) as client:
        token = entrar_admin(client, cfg)
        assert "data-sync-confirm" not in client.get("/atualizar").text
        resposta = client.post('/atualizar/sincronizar', data={
            'csrf': token, 'max_lotes': '50', 'confirmacao': 'sincronizar'})
    registros = [json.loads(line.split(' ', 2)[2]) for line in cfg.audit_log.read_text().splitlines()]
    bloqueio = next(r for r in registros if r['action'] == 'sefaz_sync')
    assert resposta.status_code == 429
    assert bloqueio['reason'] == 'fora_janela'
    assert bloqueio['proxima_tentativa'] == '2026-10-13T22:00:00-03:00'



def test_atlasdfe_usa_logos_claro_escuro_e_nome_no_cabecalho(tmp_path):
    import struct
    cfg = settings_web(tmp_path)
    criar_banco(cfg.database_path)
    app = create_app(cfg)
    with TestClient(app) as client:
        entrar_admin(client, cfg)
        resposta = client.get('/')
        for tema in ('light', 'dark'):
            url = '/static/img/atlasdfe-'+tema+'.png'
            assert url in resposta.text
            logo = client.get(url)
            assert logo.status_code == 200
            assert logo.headers['content-type'] == 'image/png'
            assert logo.content[:8] == b'\x89PNG\r\n\x1a\n'
            largura, altura = struct.unpack('>II', logo.content[16:24])
            assert largura > altura > 0
            assert logo.content[25] == 6  # PNG RGBA, com transparência.
        assert '<title>' in resposta.text and 'AtlasDFE' in resposta.text
        assert 'NexusDFE' not in resposta.text
        assert 'atlas-logo brand-logo-light' in resposta.text
        assert 'atlas-logo brand-logo-dark' in resposta.text
    assert app.title == 'AtlasDFE'



def test_admin_troca_remetente_e_testa_email_sem_sefaz(tmp_path, monkeypatch):
    from dataclasses import replace
    from nfe_consulta.web import email_alerts
    cfg = replace(settings_web(tmp_path), smtp_host='smtp.test', smtp_from='antes@gmail.com')
    criar_banco(cfg.database_path)
    app = create_app(cfg)
    enviados = []
    monkeypatch.setattr(email_alerts, '_enviar_mensagem', lambda _, mensagem: enviados.append(mensagem))
    with TestClient(app) as client:
        assert client.post('/atualizar/alertas-email/testar', data={'csrf':'x'}).status_code == 401
        token = entrar_admin(client, cfg)
        assert client.post('/atualizar/alertas-email/testar', data={'csrf':'x'}).status_code == 403
        response = client.post('/atualizar/alertas-email', data={'csrf':token,
            'recipients':'teste@empresa.com.br', 'sender':'novo@gmail.com', 'smtp_user':''}, follow_redirects=False)
        assert response.status_code == 303
        assert email_alerts.remetente_atual(cfg)[0] == 'novo@gmail.com'
        page = client.get('/atualizar').text
        assert 'Enviar e-mail de teste' in page and 'novo@gmail.com' in page
        response = client.post('/atualizar/alertas-email/testar', data={'csrf':token}, follow_redirects=False)
        assert response.status_code == 303
        assert 'email_test_sent=1' in response.headers['location']
        assert str(enviados[0]['From']) == 'novo@gmail.com'
        assert client.post('/atualizar/alertas-email/testar', data={'csrf':token}).status_code == 429
