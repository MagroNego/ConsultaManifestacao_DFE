from io import BytesIO
from pathlib import Path
import asyncio
import runpy
import sqlite3

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from conftest import authenticated_async_user, authenticated_user
from nfe_consulta.banco import BancoManifestacoes
from nfe_consulta.config import CNPJ_PADRAO
from nfe_consulta.consulta_numeros import executar_consulta_numeros, ler_numeros
from nfe_consulta.modelos import Manifestacao, RetornoDistribuicao
from nfe_consulta.web.app import create_app

web = runpy.run_path(str(Path(__file__).with_name('test_web.py')))
KEY = web['CHAVE']
NUMBER = int(KEY[25:34])


def seed(path):
    second = KEY[:22] + '002' + KEY[25:]
    no_events = KEY[:25] + str(NUMBER+1).zfill(9) + KEY[34:]
    foreign = KEY[:22] + '003' + KEY[25:]
    store = BancoManifestacoes(str(path))
    event = Manifestacao('210240','Operação não Realizada','2026-10-01T10:00:00-03:00','1')
    store.salvar_retorno(CNPJ_PADRAO, RetornoDistribuicao(138,'ok','2','2',((KEY,event),(second,event)),0))
    with store.conexao:
        store.conexao.execute('INSERT INTO informacoes_nfe(cnpj,chave,emitente,cancelada) VALUES (?,?,?,1)',(CNPJ_PADRAO,no_events,'Emitente'))
        store.conexao.execute('INSERT INTO informacoes_nfe(cnpj,chave,emitente) VALUES (?,?,?)',('outro-cnpj',foreign,'Outra empresa'))
    store.fechar()
    return second,no_events


def test_numbers_include_all_matches_missing_and_notes_without_events(tmp_path):
    path = tmp_path/'database.db'
    second,no_events = seed(path)
    before = path.read_bytes()
    source=tmp_path/'notas.txt'
    source.write_text(f'\ufeff{NUMBER}\n{NUMBER:09d}\n{NUMBER+1}\n999\n\n',encoding='utf-8')
    output=tmp_path/'report.xlsx'
    result=executar_consulta_numeros(source,path,output,CNPJ_PADRAO)
    assert {item.chave for item in result.resultados} == {KEY,second,no_events}
    assert result.total == 3 and result.com_evento == 2
    assert next(item for item in result.resultados if item.chave==no_events).cancelada
    assert path.read_bytes() == before  # Nenhuma atualização fiscal/NSU.
    book=load_workbook(output)
    assert book.active.max_row == 4
    assert {book.active.cell(row,1).value for row in (2,3,4)} == {KEY,second,no_events}
    assert all(book.active.cell(row,1).data_type=='s' for row in (2,3,4))
    summary=list(book['Numeros solicitados'].values)
    assert summary[1][:2] == (NUMBER,2)
    assert summary[2][:2] == (NUMBER+1,1)
    assert summary[3] == (999,0,'Não localizada no banco')


@pytest.mark.parametrize('content',['numero\n123','123;456','123,456','0','-1','1234567890','１２３', '1e3', '123 456'])
def test_invalid_numbers_are_rejected_before_generating(tmp_path,content):
    source=tmp_path/'notas.txt';source.write_text(content, encoding='utf-8')
    with pytest.raises(ValueError,match='Linha 1'):
        ler_numeros(source)


def test_number_limits_bom_and_empty_input(tmp_path):
    source=tmp_path/'notas.txt';source.write_text('\ufeff00123\n123\n00456\n', encoding='utf-8')
    assert list(ler_numeros(source)) == [123,456]
    with pytest.raises(ValueError):
        ler_numeros(source,max_numeros=2)
    source.write_text('\n\n')
    with pytest.raises(ValueError): ler_numeros(source)
    source.write_bytes(b'1'*(2*1024*1024+1))
    with pytest.raises(ValueError,match='2 MiB'): ler_numeros(source)


def test_local_number_export_ui_csrf_and_validation(tmp_path,monkeypatch):
    cfg=web['settings_web'](tmp_path);seed(cfg.database_path)
    def forbidden(*args,**kwargs): pytest.fail('Exportação local não deve consultar a SEFAZ')
    monkeypatch.setattr('nfe_consulta.web.app.sincronizar_configurado',forbidden)
    monkeypatch.setattr('nfe_consulta.distribuicao.subprocess.run',forbidden)
    app=create_app(cfg)
    with TestClient(app) as client:
        token=authenticated_user(client)
        html=client.get('/consulta').text
        assert 'CHAVES.txt' in html and 'NOTAS.txt' in html and 'Como preparar o arquivo TXT' in html
        args=dict(data={'csrf':token,'tipo_arquivo':'numeros'},files={'files':('notas.txt',str(NUMBER),'text/plain')})
        assert client.post('/excel',**{**args,'data':{**args['data'],'csrf':'wrong'}}).status_code==403
        response=client.post('/excel',**args)
        assert response.status_code==200
        assert 'Numeros solicitados' in load_workbook(BytesIO(response.content)).sheetnames
        invalid=client.post('/excel',data=args['data'],files={'files':('notas.txt','Número\n123','text/plain')})
        assert invalid.status_code==400 and 'Linha 1' in invalid.text
        assert client.post('/excel',**{**args,'data':{**args['data'],'tipo_arquivo':'arbitrario'}}).status_code==400


def test_number_report_survives_background_and_is_archived(tmp_path):
    async def scenario():
        cfg=web['settings_web'](tmp_path);seed(cfg.database_path)
        app=create_app(cfg)
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://test') as client:
            csrf=await authenticated_async_user(client,app)
            response=await client.post('/excel',headers={'X-NFE-Background':'1'},
                data={'csrf':csrf,'tipo_arquivo':'numeros'},files={'files':('notas.txt',str(NUMBER),'text/plain')})
            assert response.status_code==202
            identifier=response.json()['job_id']
            await asyncio.wait_for(asyncio.gather(*list(app.state.export_jobs.tasks)),5)
            item=app.state.export_jobs.records[identifier]
            assert item['status']=='complete'
            response=await client.get('/downloads/'+item['archive_id'])
            assert response.status_code==200
            book=load_workbook(BytesIO(response.content))
            assert book['Numeros solicitados']['B2'].value==2
        await app.state.export_jobs.close()
    asyncio.run(scenario())


def test_numbers_read_encrypted_database(tmp_path):
    from nfe_consulta.seguranca_banco import migrar_banco
    original=tmp_path/'plain.db'; seed(original)
    encrypted=tmp_path/'encrypted.db'
    password='Senha-SQLCipher-teste-123!'
    migrar_banco(original,encrypted,password)
    source=tmp_path/'notas.txt';source.write_text(str(NUMBER))
    result=executar_consulta_numeros(source,encrypted,tmp_path/'encrypted.xlsx',CNPJ_PADRAO,senha=password)
    assert result.total==2 and result.com_evento==2


def test_cancelled_number_export_does_not_create_report(tmp_path):
    from threading import Event
    from nfe_consulta.export_control import cancel_event, ExportCancelled
    source=tmp_path/'notas.txt';source.write_text(str(NUMBER))
    output=tmp_path/'cancelled.xlsx'
    event=Event();event.set();token=cancel_event.set(event)
    try:
        with pytest.raises(ExportCancelled):
            executar_consulta_numeros(source,tmp_path/'not-opened.db',output,CNPJ_PADRAO)
    finally:
        cancel_event.reset(token)
    assert not output.exists() and not (tmp_path/'not-opened.db').exists()
