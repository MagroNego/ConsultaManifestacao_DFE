import csv
from io import BytesIO, StringIO
from pathlib import Path
import runpy

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from nfe_consulta.banco import BancoManifestacoes
from nfe_consulta.config import CNPJ_PADRAO
from nfe_consulta.modelos import Manifestacao, RetornoDistribuicao
from nfe_consulta.seguranca_banco import abrir_banco
from nfe_consulta.web.app import create_app
from nfe_consulta.web import xml_store
from nfe_consulta.web.consulta_local import consultar_eventos, normalizar_filtros

helpers = runpy.run_path(str(Path(__file__).with_name('test_cancelamento_xml.py')))
KEY = helpers['KEY']
KEY2 = KEY[:25] + '000000124' + KEY[34:]


def setup(tmp_path):
    cfg = helpers['web']['settings_web'](tmp_path)
    banco = BancoManifestacoes(str(cfg.database_path))
    banco.salvar_retorno(CNPJ_PADRAO, RetornoDistribuicao(138, 'ok', '2', '2', tuple(
        (key, Manifestacao('210240', 'Operação não Realizada', '2026-10-01T10:00:00-03:00', '135260000000001', str(index)))
        for index, key in enumerate((KEY, KEY2), 1))))
    banco.fechar()
    helpers['batch'](cfg, tmp_path, [helpers['helpers']['xml'](), helpers['helpers']['xml'](chave=KEY2), helpers['cancellation']()])
    return cfg


@pytest.mark.parametrize('state,key', [('cancelada', KEY), ('autorizada', KEY2)])
def test_filters_apply_to_all_reports_and_exports(tmp_path, state, key):
    cfg = setup(tmp_path)
    expected = state.capitalize()
    for kind in ('notas', 'itens', 'retencoes'):
        report = xml_store.query_report(cfg.database_path, CNPJ_PADRAO, kind=kind, situacao=state)
        assert report['total'] == 1
        assert report['rows'][0]['Chave Acesso'] == key
        assert report['rows'][0]['Situação'] == expected
    events = consultar_eventos(cfg.database_path, CNPJ_PADRAO, normalizar_filtros(situacao=state))
    assert events.total == 1 and events.eventos[0].chave == key
    with TestClient(create_app(cfg)) as client:
        for url in ('/xml', '/consulta'):
            page = client.get(url, params={'situacao': state}).text
            assert f'<option value="{state}" selected>' in page
            assert f'name="situacao" value="{state}"' in page
            rows = page.split('<tbody>')[1].split('</tbody>')[0]
            assert key in rows
            assert (KEY2 if key == KEY else KEY) not in rows
        for kind in ('notas', 'itens', 'retencoes'):
            response = client.get('/xml/exportar', params={'formato': 'csv', 'tipo': kind, 'situacao': state})
            rows = list(csv.DictReader(StringIO(response.content.decode('utf-8-sig')), delimiter=';'))
            assert len(rows) == 1 and rows[0]['Chave Acesso'] == key
            assert rows[0]['Situação'] == expected
        for url, sheets, header_row in [('/xml/exportar', ('Notas', 'Itens', 'Retencoes'), 1), ('/consulta/exportar', ('Manifestacoes',), 4)]:
            response = client.get(url, params={'situacao': state, 'formato': 'xlsx'})
            assert response.status_code == 200
            wb = load_workbook(BytesIO(response.content))
            for sheet in sheets:
                ws = wb[sheet]
                headers = [cell.value for cell in ws[header_row]]
                assert ws.max_row == header_row + 1
                assert ws.cell(header_row + 1, headers.index('Situação') + 1).value == expected


def test_state_filter_before_count_and_pagination(tmp_path):
    cfg = helpers['web']['settings_web'](tmp_path)
    helpers['web']['criar_banco'](cfg.database_path)
    notes = [helpers['helpers']['xml'](chave=KEY[:25] + f'{number:09d}' + KEY[34:]) for number in range(123, 225)]
    helpers['batch'](cfg, tmp_path, [*notes, helpers['cancellation']()])
    report = xml_store.query_report(cfg.database_path, CNPJ_PADRAO, situacao='autorizada', page=2)
    assert report['total'] == 101 and report['pages'] == 2 and len(report['rows']) == 1
    report = xml_store.query_report(cfg.database_path, CNPJ_PADRAO, situacao='cancelada', page=2)
    assert report['total'] == 1 and report['page'] == report['pages'] == 1
    with TestClient(create_app(cfg)) as client:
        page = client.get('/xml?pagina=1&situacao=autorizada').text
        assert 'situacao=autorizada' in page and 'pagina=2' in page


def test_invalid_filter_and_legacy_database_without_status_table(tmp_path):
    cfg = setup(tmp_path)
    with pytest.raises(ValueError):
        normalizar_filtros(situacao='invalid')
    with pytest.raises(ValueError):
        xml_store.query_report(cfg.database_path, CNPJ_PADRAO, situacao='invalid')
    with TestClient(create_app(cfg)) as client:
        assert client.get('/xml/exportar?formato=csv&situacao=invalid').status_code == 400
        assert 'Situação inválida' in client.get('/consulta?situacao=invalid').text
    conn = abrir_banco(cfg.database_path, None)
    conn.execute('DROP TABLE informacoes_nfe')
    conn.commit()
    conn.close()
    assert xml_store.query_report(cfg.database_path, CNPJ_PADRAO, situacao='cancelada')['total'] == 0
    assert xml_store.query_report(cfg.database_path, CNPJ_PADRAO, situacao='autorizada')['total'] == 2
    assert consultar_eventos(cfg.database_path, CNPJ_PADRAO, normalizar_filtros(situacao='cancelada')).total == 0
